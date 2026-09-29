"""`read_excel` and `ExcelFile`, a port of pandas' Excel reader.

pandas reads a workbook in two steps. An engine (openpyxl, odfpy, xlrd, pyxlsb
or python-calamine) turns a sheet into rows of Python values, each cell already
typed as a number, a flag, a moment or text. Then the rows go through the
parser behind `read_csv(engine="python")`, which picks the header, the names,
the row labels and the columns kept, and types each column from the values it
holds. This module ports both steps. The engines are pandas' own few lines per
package, and the parser is ported from pandas' Python parser, working on lists
where pandas works on numpy object arrays, so the same cells give the same
columns: a whole number that Excel stores as a float reads as an integer, a
column of flags with a gap reads as floats, and text that looks like a number
reads as one.
"""

from __future__ import annotations

import contextlib
import datetime as dt
import math
import numbers
import os
import re
import warnings
import zipfile
from collections import defaultdict
from collections.abc import Callable, Hashable, Iterable, Sequence
from typing import Any

from . import _objects, _optional
from ._pandas import NO_DEFAULT
from ._textread import DEFAULT_NA
from .errors import EmptyDataError, InvalidArgumentError, ParserError, ParserWarning

__all__ = ["ExcelFile", "read_excel"]

_NAN = float("nan")

_XLS_SIGNATURES = (
    b"\x09\x00\x04\x00\x07\x00\x10\x00",
    b"\x09\x02\x06\x00\x00\x00\x10\x00",
    b"\x09\x04\x06\x00\x00\x00\x10\x00",
    b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1",
)
_ZIP_SIGNATURE = b"PK\x03\x04"
_PEEK_SIZE = max(map(len, (*_XLS_SIGNATURES, _ZIP_SIGNATURE)))

_DEFAULT_READERS = {
    "xlsx": "openpyxl",
    "xlsm": "openpyxl",
    "xlsb": "pyxlsb",
    "xls": "xlrd",
    "ods": "odf",
}

_STORAGE = "storage_options passed with file object or non-fsspec file path"
_USECOLS = (
    "'usecols' must either be list-like of all strings, all unicode, all integers or a callable."
)
_TRUE = frozenset({"True", "TRUE", "true"})
_FALSE = frozenset({"False", "FALSE", "false"})
_INT64 = (-(1 << 63), (1 << 63) - 1)
_UINT64 = (1 << 64) - 1
# What C's isspace skips, which is what pandas' number parser skips around a number.
_SPACE = " \t\n\v\f\r"
_NUMBER = re.compile(r"[+-]?(?:([0-9]+)(\.[0-9]*)?|(\.[0-9]+))(?:([eE])[+-]?[0-9]+)?")
_INFINITY = frozenset({"inf", "-inf", "+inf", "infinity", "-infinity", "+infinity"})


def _is_integer(value: Any) -> bool:
    return isinstance(value, numbers.Integral) and not isinstance(value, bool)


def _is_float(value: Any) -> bool:
    return isinstance(value, float) or type(value).__name__ in ("float32", "float16", "float64")


def _is_list_like(value: Any, allow_sets: bool = True) -> bool:
    if isinstance(value, (str, bytes)) or not hasattr(value, "__iter__"):
        return False
    return allow_sets or not isinstance(value, (set, frozenset))


def _missing(value: Any) -> bool:
    return value is None or (isinstance(value, float) and value != value)


def _validate_integer(name: str, value: Any) -> Any:
    if value is None:
        return value
    message = f"'{name}' must be an integer >=0"
    if _is_float(value):
        if int(value) != value:
            raise InvalidArgumentError(message)
        value = int(value)
    elif not (_is_integer(value) and value >= 0):
        raise InvalidArgumentError(message)
    return int(value)


def _validate_header(header: Any) -> None:
    if header is None:
        return
    if _is_integer(header):
        if header < 0:
            raise InvalidArgumentError(
                "Passing negative integer to header is invalid. For no header, use header=None"
                " instead"
            )
        return
    if _is_list_like(header, allow_sets=False):
        if not all(map(_is_integer, header)):
            raise InvalidArgumentError("header must be integer or list of integers")
        if any(i < 0 for i in header):
            raise InvalidArgumentError("cannot specify multi-index header with negative integers")
        return
    if isinstance(header, bool):
        raise TypeError(
            "Passing a bool to header is invalid. Use header=None for no header or header=int"
            " or list-like of ints to specify the row(s) making up the column names"
        )
    raise InvalidArgumentError("header must be integer or list of integers")


# The engines. Each turns a sheet into rows of Python values, as pandas' does.


def _opened(source: Any, storage_options: Any) -> tuple[Any, bool]:
    """A binary handle on a path, a URL or a handle, and whether it was opened here."""
    if hasattr(source, "read"):
        if storage_options is not None:
            raise InvalidArgumentError(_STORAGE)
        return source, False
    if isinstance(source, (bytes, bytearray)):
        raise TypeError(f"Expected file path name or file-like object, got {type(source)} type")
    path = os.fspath(source)
    if isinstance(path, bytes):
        path = path.decode()
    if "://" in path:
        import io

        from ._textread import _fetched

        return io.BytesIO(_fetched(path, storage_options)), True
    if storage_options is not None:
        raise InvalidArgumentError(_STORAGE)
    return open(os.path.expanduser(path), "rb"), True


class _Reader:
    """What pandas' `BaseExcelReader` does for every engine."""

    module = ""
    extra = ""

    def __init__(self, source: Any, storage_options: Any, engine_kwargs: Any) -> None:
        _optional.imported(self.module, self.extra)
        engine_kwargs = {} if engine_kwargs is None else engine_kwargs
        self.handle = None
        if isinstance(source, self.workbook_class()):
            self.book = source
            return
        handle, owned = _opened(source, storage_options)
        if owned:
            self.handle = handle
        handle.seek(0)
        try:
            self.book = self.load_workbook(handle, engine_kwargs)
        except Exception:
            self.close()
            raise

    def workbook_class(self) -> Any:
        raise NotImplementedError

    def load_workbook(self, handle: Any, engine_kwargs: dict[str, Any]) -> Any:
        raise NotImplementedError

    @property
    def sheet_names(self) -> list[str]:
        raise NotImplementedError

    def sheet_by_name(self, name: str) -> Any:
        raise NotImplementedError

    def sheet_by_index(self, index: int) -> Any:
        raise NotImplementedError

    def sheet_data(self, sheet: Any, rows: int | None) -> list[list[Any]]:
        raise NotImplementedError

    def close(self) -> None:
        book = getattr(self, "book", None)
        if book is not None:
            if hasattr(book, "close"):
                book.close()
            elif hasattr(book, "release_resources"):
                book.release_resources()
        if self.handle is not None:
            self.handle.close()

    def check_index(self, index: int) -> None:
        count = len(self.sheet_names)
        if index >= count:
            raise InvalidArgumentError(
                f"Worksheet index {index} is invalid, {count} worksheets found"
            )

    def check_name(self, name: str) -> None:
        if name not in self.sheet_names:
            raise InvalidArgumentError(f"Worksheet named '{name}' not found")


def _padded(data: list[list[Any]]) -> list[list[Any]]:
    if data:
        width = max(len(row) for row in data)
        if min(len(row) for row in data) < width:
            data = [row + [""] * (width - len(row)) for row in data]
    return data


def _whole(value: float) -> int | float:
    number = int(value)
    return number if number == value else float(value)


class _Openpyxl(_Reader):
    module = "openpyxl"

    def workbook_class(self) -> Any:
        from openpyxl import Workbook

        return Workbook

    def load_workbook(self, handle: Any, engine_kwargs: dict[str, Any]) -> Any:
        from openpyxl import load_workbook

        defaults = {"read_only": True, "data_only": True, "keep_links": False}
        return load_workbook(handle, **defaults | engine_kwargs)

    @property
    def sheet_names(self) -> list[str]:
        return [sheet.title for sheet in self.book.worksheets]

    def sheet_by_name(self, name: str) -> Any:
        self.check_name(name)
        return self.book[name]

    def sheet_by_index(self, index: int) -> Any:
        self.check_index(index)
        return self.book.worksheets[index]

    @staticmethod
    def _cell(cell: Any) -> Any:
        from openpyxl.cell.cell import TYPE_ERROR, TYPE_NUMERIC

        if cell.value is None:
            return ""
        if cell.data_type == TYPE_ERROR:
            return _NAN
        if cell.data_type == TYPE_NUMERIC:
            return _whole(cell.value)
        return cell.value

    def sheet_data(self, sheet: Any, rows: int | None) -> list[list[Any]]:
        if self.book.read_only:
            sheet.reset_dimensions()
        data: list[list[Any]] = []
        last = -1
        for number, row in enumerate(sheet.rows):
            converted = [self._cell(cell) for cell in row]
            while converted and converted[-1] == "":
                converted.pop()
            if converted:
                last = number
            data.append(converted)
            if rows is not None and len(data) >= rows:
                break
        return _padded(data[: last + 1])


class _Odf(_Reader):
    module = "odf"

    def workbook_class(self) -> Any:
        from odf.opendocument import OpenDocument

        return OpenDocument

    def load_workbook(self, handle: Any, engine_kwargs: dict[str, Any]) -> Any:
        from odf.opendocument import load

        return load(handle, **engine_kwargs)

    def _tables(self) -> list[Any]:
        from odf.table import Table

        return self.book.getElementsByType(Table)

    @property
    def sheet_names(self) -> list[str]:
        return [table.getAttribute("name") for table in self._tables()]

    def sheet_by_index(self, index: int) -> Any:
        self.check_index(index)
        return self._tables()[index]

    def sheet_by_name(self, name: str) -> Any:
        self.check_name(name)
        for table in self._tables():
            if table.getAttribute("name") == name:
                return table
        self.close()
        raise InvalidArgumentError(f"sheet {name} not found")

    def sheet_data(self, sheet: Any, rows: int | None) -> list[list[Any]]:
        from odf.namespaces import TABLENS
        from odf.table import CoveredTableCell, TableCell, TableRow

        covered = CoveredTableCell().qname
        plain = TableCell().qname
        names = {covered, plain}
        empty_rows = 0
        widest = 0
        table: list[list[Any]] = []
        for sheet_row in sheet.getElementsByType(TableRow):
            empty_cells = 0
            row: list[Any] = []
            for cell in sheet_row.childNodes:
                if hasattr(cell, "qname") and cell.qname in names:
                    value = self._value(cell) if cell.qname == plain else ""
                    repeat = int(cell.attributes.get((TABLENS, "number-columns-repeated"), 1))
                    if isinstance(value, str) and value == "":
                        empty_cells += repeat
                    else:
                        row.extend([""] * empty_cells)
                        empty_cells = 0
                        row.extend([value] * repeat)
            widest = max(widest, len(row))
            repeat = int(sheet_row.attributes.get((TABLENS, "number-rows-repeated"), 1))
            if not row:
                empty_rows += repeat
            else:
                table.extend([[""]] * empty_rows)
                empty_rows = 0
                table.extend(row for _ in range(repeat))
            if rows is not None and len(table) >= rows:
                break
        for row in table:
            if len(row) < widest:
                row.extend([""] * (widest - len(row)))
        return table

    def _value(self, cell: Any) -> Any:
        from odf.namespaces import OFFICENS

        from ._scalars import Timestamp

        if str(cell) == "#N/A":
            return _NAN
        kind = cell.attributes.get((OFFICENS, "value-type"))
        if kind == "boolean":
            return str(cell) == "TRUE"
        if kind is None:
            return ""
        if kind == "float":
            return _whole(float(cell.attributes.get((OFFICENS, "value"))))
        if kind in ("percentage", "currency"):
            return float(cell.attributes.get((OFFICENS, "value")))
        if kind == "string":
            return self._text(cell)
        if kind == "date":
            return Timestamp(cell.attributes.get((OFFICENS, "date-value")))
        if kind == "time":
            return Timestamp(str(cell)).time()
        self.close()
        raise InvalidArgumentError(f"Unrecognized type {kind}")

    def _text(self, cell: Any) -> str:
        from odf.element import Element
        from odf.namespaces import TEXTNS
        from odf.office import Annotation
        from odf.text import S

        annotation = Annotation().qname
        space = S().qname
        out = []
        for fragment in cell.childNodes:
            if isinstance(fragment, Element):
                if fragment.qname == space:
                    out.append(" " * int(fragment.attributes.get((TEXTNS, "c"), 1)))
                elif fragment.qname == annotation:
                    continue
                else:
                    out.append(self._text(fragment))
            else:
                out.append(str(fragment).strip("\n"))
        return "".join(out)


class _Xlrd(_Reader):
    module = "xlrd"
    extra = "Install xlrd >= 2.0.1 for xls Excel support"

    def workbook_class(self) -> Any:
        from xlrd import Book

        return Book

    def load_workbook(self, handle: Any, engine_kwargs: dict[str, Any]) -> Any:
        from xlrd import open_workbook

        return open_workbook(file_contents=handle.read(), **engine_kwargs)

    @property
    def sheet_names(self) -> list[str]:
        return self.book.sheet_names()

    def sheet_by_name(self, name: str) -> Any:
        self.check_name(name)
        return self.book.sheet_by_name(name)

    def sheet_by_index(self, index: int) -> Any:
        self.check_index(index)
        return self.book.sheet_by_index(index)

    def sheet_data(self, sheet: Any, rows: int | None) -> list[list[Any]]:
        from xlrd import XL_CELL_BOOLEAN, XL_CELL_DATE, XL_CELL_ERROR, XL_CELL_NUMBER, xldate

        epoch1904 = self.book.datemode

        def parsed(value: Any, kind: int) -> Any:
            if kind == XL_CELL_DATE:
                try:
                    value = xldate.xldate_as_datetime(value, epoch1904)
                except OverflowError:
                    return value
                year = value.timetuple()[0:3]
                if (not epoch1904 and year == (1899, 12, 31)) or (
                    epoch1904 and year == (1904, 1, 1)
                ):
                    value = dt.time(value.hour, value.minute, value.second, value.microsecond)
            elif kind == XL_CELL_ERROR:
                value = _NAN
            elif kind == XL_CELL_BOOLEAN:
                value = bool(value)
            elif kind == XL_CELL_NUMBER and math.isfinite(value):
                number = int(value)
                if number == value:
                    value = number
            return value

        count = sheet.nrows if rows is None else min(sheet.nrows, rows)
        return [
            [
                parsed(value, kind)
                for value, kind in zip(sheet.row_values(i), sheet.row_types(i), strict=True)
            ]
            for i in range(count)
        ]


class _Pyxlsb(_Reader):
    module = "pyxlsb"

    def workbook_class(self) -> Any:
        from pyxlsb import Workbook

        return Workbook

    def load_workbook(self, handle: Any, engine_kwargs: dict[str, Any]) -> Any:
        from pyxlsb import open_workbook

        return open_workbook(handle, **engine_kwargs)

    @property
    def sheet_names(self) -> list[str]:
        return self.book.sheets

    def sheet_by_name(self, name: str) -> Any:
        self.check_name(name)
        return self.book.get_sheet(name)

    def sheet_by_index(self, index: int) -> Any:
        self.check_index(index)
        return self.book.get_sheet(index + 1)

    def sheet_data(self, sheet: Any, rows: int | None) -> list[list[Any]]:
        data: list[list[Any]] = []
        previous = -1
        for row in sheet.rows(sparse=True):
            number = row[0].r
            converted = [
                "" if cell.v is None else _whole(cell.v) if isinstance(cell.v, float) else cell.v
                for cell in row
            ]
            while converted and converted[-1] == "":
                converted.pop()
            if converted:
                data.extend([[]] * (number - previous - 1))
                data.append(converted)
                previous = number
            if rows is not None and len(data) >= rows:
                break
        return _padded(data)


class _Calamine(_Reader):
    module = "python_calamine"

    def workbook_class(self) -> Any:
        from python_calamine import CalamineWorkbook

        return CalamineWorkbook

    def load_workbook(self, handle: Any, engine_kwargs: dict[str, Any]) -> Any:
        from python_calamine import load_workbook

        return load_workbook(handle, **engine_kwargs)

    @property
    def sheet_names(self) -> list[str]:
        from python_calamine import SheetTypeEnum

        return [
            sheet.name
            for sheet in self.book.sheets_metadata
            if sheet.typ == SheetTypeEnum.WorkSheet
        ]

    def sheet_by_name(self, name: str) -> Any:
        self.check_name(name)
        return self.book.get_sheet_by_name(name)

    def sheet_by_index(self, index: int) -> Any:
        self.check_index(index)
        return self.book.get_sheet_by_index(index)

    def sheet_data(self, sheet: Any, rows: int | None) -> list[list[Any]]:
        def converted(value: Any) -> Any:
            if isinstance(value, float):
                return _whole(value)
            if isinstance(value, (dt.datetime, dt.timedelta, dt.time)):
                return value
            if isinstance(value, dt.date):
                return dt.datetime(value.year, value.month, value.day)
            return value

        found = sheet.to_python(skip_empty_area=False, nrows=rows)
        return [[converted(cell) for cell in row] for row in found]


_ENGINES: dict[str, type[_Reader]] = {
    "xlrd": _Xlrd,
    "openpyxl": _Openpyxl,
    "odf": _Odf,
    "pyxlsb": _Pyxlsb,
    "calamine": _Calamine,
}


def _excel_format(source: Any, storage_options: Any) -> str | None:
    """pandas' `inspect_excel_format`: the kind of workbook, from its first bytes."""
    handle, owned = _opened(source, storage_options)
    try:
        handle.seek(0)
        peek = handle.read(_PEEK_SIZE)
        if peek is None:
            raise InvalidArgumentError("stream is empty")
        handle.seek(0)
        if any(peek.startswith(signature) for signature in _XLS_SIGNATURES):
            return "xls"
        if not peek.startswith(_ZIP_SIGNATURE):
            return None
        with zipfile.ZipFile(handle) as archive:
            names = {name.replace("\\", "/").lower() for name in archive.namelist()}
        handle.seek(0)
        if "xl/workbook.xml" in names:
            return "xlsx"
        if "xl/workbook.bin" in names:
            return "xlsb"
        if "content.xml" in names:
            return "ods"
        return "zip"
    finally:
        if owned:
            handle.close()


class ExcelFile:
    """A workbook opened once, to read several sheets from, as pandas' `ExcelFile`.

    `engine` is one of pandas' five, picked from the file's first bytes when not
    given: openpyxl for xlsx and xlsm, pyxlsb for xlsb, xlrd for xls and odf
    for ods, unless the `io.excel.<kind>.reader` option names another. The
    package behind the engine has to be installed, and a missing one is named
    in pandas' words.
    """

    _engines = _ENGINES

    def __init__(
        self,
        path_or_buffer: Any,
        engine: str | None = None,
        storage_options: Any = None,
        engine_kwargs: dict | None = None,
    ) -> None:
        from ._config import get_option

        if engine_kwargs is None:
            engine_kwargs = {}
        if engine is not None and engine not in self._engines:
            raise InvalidArgumentError(f"Unknown engine: {engine}")
        self._io = (
            os.fspath(path_or_buffer) if isinstance(path_or_buffer, os.PathLike) else path_or_buffer
        )
        if engine is None:
            kind = _excel_format(path_or_buffer, storage_options)
            if kind is None:
                raise InvalidArgumentError(
                    "Excel file format cannot be determined, you must specify an engine manually."
                )
            engine = get_option(f"io.excel.{kind}.reader")
            if engine == "auto":
                engine = _DEFAULT_READERS[kind]
        self.engine = engine
        self.storage_options = storage_options
        self._reader = self._engines[engine](
            self._io, storage_options=storage_options, engine_kwargs=engine_kwargs
        )

    def __fspath__(self) -> Any:
        return self._io

    def parse(
        self,
        sheet_name: str | int | list[int] | list[str] | None = 0,
        header: int | Sequence[int] | None = 0,
        names: Any = None,
        index_col: int | Sequence[int] | None = None,
        usecols: Any = None,
        converters: Any = None,
        true_values: Iterable[Hashable] | None = None,
        false_values: Iterable[Hashable] | None = None,
        skiprows: Sequence[int] | int | Callable[[int], object] | None = None,
        nrows: int | None = None,
        na_values: Any = None,
        parse_dates: list | dict | bool = False,
        date_format: str | dict[Hashable, str] | None = None,
        thousands: str | None = None,
        comment: str | None = None,
        skipfooter: int = 0,
        dtype_backend: Any = NO_DEFAULT,
        **kwds: Any,
    ) -> Any:
        """One sheet as a frame, or a dict of frames for a list of sheets or `None`.

        The arguments are `read_excel`'s, and are read the same way.
        """
        return _parse(
            self._reader,
            sheet_name=sheet_name,
            header=header,
            names=names,
            index_col=index_col,
            usecols=usecols,
            converters=converters,
            true_values=true_values,
            false_values=false_values,
            skiprows=skiprows,
            nrows=nrows,
            na_values=na_values,
            parse_dates=parse_dates,
            date_format=date_format,
            thousands=thousands,
            comment=comment,
            skipfooter=skipfooter,
            dtype_backend=dtype_backend,
            **kwds,
        )

    @property
    def book(self) -> Any:
        """The engine's own workbook object."""
        return self._reader.book

    @property
    def sheet_names(self) -> list[str]:
        """The names of the sheets, in the workbook's order."""
        return self._reader.sheet_names

    def close(self) -> None:
        """Closes the workbook and the file under it."""
        self._reader.close()

    def __enter__(self) -> ExcelFile:
        return self

    def __exit__(self, *exc: object) -> None:
        self.close()


def read_excel(
    io: Any,
    sheet_name: str | int | list | None = 0,
    *,
    header: int | Sequence[int] | None = 0,
    names: Any = None,
    index_col: int | str | Sequence[int] | None = None,
    usecols: Any = None,
    dtype: Any = None,
    engine: str | None = None,
    converters: Any = None,
    true_values: Iterable[Hashable] | None = None,
    false_values: Iterable[Hashable] | None = None,
    skiprows: Sequence[int] | int | Callable[[int], object] | None = None,
    nrows: int | None = None,
    na_values: Any = None,
    keep_default_na: bool = True,
    na_filter: bool = True,
    verbose: bool = False,
    parse_dates: list | dict | bool = False,
    date_format: dict[Hashable, str] | str | None = None,
    thousands: str | None = None,
    decimal: str = ".",
    comment: str | None = None,
    skipfooter: int = 0,
    storage_options: Any = None,
    dtype_backend: Any = NO_DEFAULT,
    engine_kwargs: dict | None = None,
) -> Any:
    """A sheet of an Excel or OpenDocument workbook as a frame, as pandas reads it.

    `io` is a path, a URL, a binary handle or an `ExcelFile`. `sheet_name` is a
    position, a name, a list of either, or `None` for every sheet, and a list or
    `None` answers a dict of frames keyed the way the sheets were asked for.
    The rest are `read_csv`'s arguments, read by pandas' Python parser on the
    typed cells, so a number stays a number, a moment a moment, and text that
    looks like a number reads as one. A list-like `index_col` fills a gap in a
    label column from the label above it, as pandas does for merged cells.

    Raises:
        ImportError: When the engine's package is not installed, in pandas' words.
        ValueError: For the mistakes pandas names, in its words, each ending
            with the sheet it was reading.
    """
    from ._pandas import _backend

    if dtype_backend is not NO_DEFAULT:
        _backend(dtype_backend)
    close = False
    if engine_kwargs is None:
        engine_kwargs = {}
    if not isinstance(io, ExcelFile):
        close = True
        io = ExcelFile(
            io, storage_options=storage_options, engine=engine, engine_kwargs=engine_kwargs
        )
    elif engine and engine != io.engine:
        raise InvalidArgumentError(
            "Engine should not be specified when passing an ExcelFile - ExcelFile already has"
            " the engine set"
        )
    try:
        return io.parse(
            sheet_name=sheet_name,
            header=header,
            names=names,
            index_col=index_col,
            usecols=usecols,
            dtype=dtype,
            converters=converters,
            true_values=true_values,
            false_values=false_values,
            skiprows=skiprows,
            nrows=nrows,
            na_values=na_values,
            keep_default_na=keep_default_na,
            na_filter=na_filter,
            verbose=verbose,
            parse_dates=parse_dates,
            date_format=date_format,
            thousands=thousands,
            decimal=decimal,
            comment=comment,
            skipfooter=skipfooter,
            dtype_backend=dtype_backend,
        )
    finally:
        if close:
            io.close()


# pandas' `BaseExcelReader.parse` and `_parse_sheet`.


def _rows_needed(skiprows: Any, wanted: int) -> int:
    at = 0
    used = 0
    while used < wanted:
        if not skiprows(at):
            used += 1
        at += 1
    return at


def _rows_to_read(header: Any, index_col: Any, skiprows: Any, nrows: Any) -> int | None:
    """How many rows of the sheet pandas reads for `nrows`, so a big sheet is not read whole."""
    if nrows is None:
        return None
    if header is None:
        header_rows = 1
    elif _is_integer(header):
        header_rows = 1 + header
    else:
        header_rows = 1 + header[-1]
    if _is_list_like(header) and index_col is not None and len(header) > 1:
        header_rows += 1
    if skiprows is None:
        return header_rows + nrows
    if _is_integer(skiprows):
        return header_rows + nrows + skiprows
    if _is_list_like(skiprows):
        listed = skiprows
        return _rows_needed(lambda x: x in listed, header_rows + nrows)
    if callable(skiprows):
        return _rows_needed(skiprows, header_rows + nrows)
    return None


def _column_number(letters: str) -> int:
    index = 0
    for c in letters.upper().strip():
        code = ord(c)
        if code < ord("A") or code > ord("Z"):
            raise InvalidArgumentError(f"Invalid column name: {letters}")
        index = index * 26 + code - ord("A") + 1
    return index - 1


def _usecols(usecols: Any) -> Any:
    """pandas' `maybe_convert_usecols`: Excel's letters and ranges as positions."""
    if usecols is None:
        return usecols
    if _is_integer(usecols):
        raise InvalidArgumentError(
            "Passing an integer for `usecols` is no longer supported.  Please pass in a list of"
            " int from 0 to `usecols` inclusive instead."
        )
    if isinstance(usecols, str):
        out: list[int] = []
        for area in usecols.split(","):
            if ":" in area:
                ends = area.split(":")
                out.extend(range(_column_number(ends[0]), _column_number(ends[1]) + 1))
            else:
                out.append(_column_number(area))
        return out
    return usecols


def _parse(reader: _Reader, sheet_name: Any = 0, **options: Any) -> Any:
    from ._frame import DataFrame

    header = options.get("header", 0)
    _validate_header(header)
    _validate_integer("nrows", options.get("nrows"))
    many = False
    if isinstance(sheet_name, list):
        sheets = sheet_name
        many = True
    elif sheet_name is None:
        sheets = reader.sheet_names
        many = True
    else:
        sheets = [sheet_name]
    sheets = list(dict.fromkeys(sheets))
    output: dict[Any, Any] = {}
    last = None
    for name in sheets:
        last = name
        if options.get("verbose"):
            print(f"Reading sheet {name}")
        sheet = reader.sheet_by_name(name) if isinstance(name, str) else reader.sheet_by_index(name)
        needed = _rows_to_read(
            header, options.get("index_col"), options.get("skiprows"), options.get("nrows")
        )
        data = reader.sheet_data(sheet, needed)
        if hasattr(sheet, "close"):
            sheet.close()
        options["usecols"] = _usecols(options.get("usecols"))
        if not data:
            output[name] = DataFrame()
            continue
        output[name] = _parse_sheet(data, name, dict(options))
    if last is None:
        raise InvalidArgumentError("Sheet name is an empty list")
    return output if many else output[last]


def _fill_header(row: list[Any], control: list[bool]) -> tuple[list[Any], list[bool]]:
    last = row[0]
    for i in range(1, len(row)):
        if not control[i]:
            last = row[i]
        if _blank_cell(row[i]):
            row[i] = last
        else:
            control[i] = False
            last = row[i]
    return row, control


def _blank_cell(value: Any) -> bool:
    return value is None or (isinstance(value, str) and value == "")


def _parse_sheet(data: list[list[Any]], sheet: Any, options: dict[str, Any]) -> Any:
    from ._frame import DataFrame

    header = options.get("header", 0)
    index_col = options.get("index_col")
    skiprows = options.get("skiprows")
    listed_header = _is_list_like(header)
    single = listed_header and len(header) == 1
    if single:
        header = header[0]
        options["header"] = header
    header_names = None
    control: list[bool] = []
    if header is not None and _is_list_like(header):
        header_names = []
        control = [True] * len(data[0])
        for row in header:
            if _is_integer(skiprows):
                row += skiprows
            if row > len(data) - 1:
                raise InvalidArgumentError(
                    f"header index {row} exceeds maximum index {len(data) - 1} of data."
                )
            data[row], control = _fill_header(data[row], control)
            if index_col is not None:
                at = max(index_col) if _is_list_like(index_col) else index_col
                found = data[row][at]
                header_names.append(None if _blank_cell(found) else found)
    has_index_names = False
    if listed_header and not single and index_col is not None:
        wanted = {index_col} if isinstance(index_col, int) else set(index_col)
        if len(header) < len(data):
            potential = data[len(header)]
            has_index_names = all(
                _blank_cell(x)
                for i, x in enumerate(potential)
                if not control[i] and i not in wanted
            )
    if _is_list_like(index_col):
        if header is None:
            offset = 0
        elif isinstance(header, int):
            offset = 1 + header
        else:
            offset = 1 + max(header)
        if has_index_names:
            offset += 1
        if offset < len(data):
            for col in index_col:
                last = data[offset][col]
                for row in range(offset + 1, len(data)):
                    if _blank_cell(data[row][col]):
                        data[row][col] = last
                    else:
                        last = data[row][col]
    options["has_index_names"] = has_index_names
    try:
        frame = _TextParser(data, options).read(options.get("nrows"))
        if header_names:
            frame = frame.set_axis(frame.columns.set_names(header_names), axis=1)
        return frame
    except EmptyDataError:
        return DataFrame()
    except Exception as err:
        if err.args:
            err.args = (f"{err.args[0]} (sheet: {sheet})", *err.args[1:])
        raise


# pandas' Python parser, on rows of Python values.


def _stringified_na(values: Iterable[Any]) -> set[Any]:
    out: list[Any] = []
    for x in values:
        out.append(str(x))
        out.append(x)
        try:
            v = float(x)
            if v == int(v):
                v = int(v)
                out.append(f"{v}.0")
                out.append(str(v))
            out.append(v)
        except (TypeError, ValueError, OverflowError):
            pass
        with contextlib.suppress(TypeError, ValueError, OverflowError):
            out.append(int(x))
    return set(out)


def _floated_na(values: Iterable[Any]) -> set[float]:
    out = set()
    for v in values:
        try:
            number = float(v)
        except (TypeError, ValueError, OverflowError):
            continue
        if not math.isnan(number):
            out.add(number)
    return out


def _clean_na(na_values: Any, keep_default_na: bool) -> tuple[Any, Any]:
    if na_values is None:
        return (set(DEFAULT_NA) if keep_default_na else set()), set()
    if isinstance(na_values, dict):
        cleaned = {}
        for key, value in na_values.items():
            if not _is_list_like(value):
                value = [value]
            if keep_default_na:
                value = set(value) | DEFAULT_NA
            cleaned[key] = _stringified_na(value)
        return cleaned, {key: _floated_na(value) for key, value in cleaned.items()}
    if not _is_list_like(na_values):
        na_values = [na_values]
    cleaned_set = _stringified_na(na_values)
    if keep_default_na:
        cleaned_set = cleaned_set | DEFAULT_NA
    return cleaned_set, _floated_na(cleaned_set)


def _na_for(column: Any, na_values: Any, na_fvalues: Any, keep_default_na: bool) -> tuple:
    if isinstance(na_values, dict):
        if column in na_values:
            return na_values[column], na_fvalues[column]
        if keep_default_na:
            return set(DEFAULT_NA), set()
        return set(), set()
    return na_values, na_fvalues


def _in(value: Any, values: Any) -> bool:
    try:
        return value in values
    except TypeError:
        return False


def _dedup(names: Sequence[Any], tuples: bool) -> list[Any]:
    names = list(names)
    counts: defaultdict[Any, int] = defaultdict(int)
    for i, col in enumerate(names):
        count = counts[col]
        while count > 0:
            counts[col] = count + 1
            col = (*col[:-1], f"{col[-1]}.{count}") if tuples else f"{col}.{count}"
            count = counts[col]
        names[i] = col
        counts[col] = count + 1
    return names


def _potential_multi(columns: Sequence[Any], index_col: Any = None) -> bool:
    skipped = set() if index_col is None or isinstance(index_col, bool) else set(index_col)
    return bool(len(columns)) and all(isinstance(c, tuple) for c in columns if not _in(c, skipped))


class _Values:
    """A parsed column: its values and the numpy type pandas would hold them as."""

    __slots__ = ("kind", "mask", "values")

    def __init__(self, kind: str, values: list[Any], mask: list[bool] | None = None) -> None:
        self.kind = kind
        self.values = values
        self.mask = mask


def _floatify(text: Any) -> tuple[float, bool]:
    """pandas' `floatify`: a number from text, and whether it could be an integer."""
    if isinstance(text, bytes):
        text = text.decode()
    if not isinstance(text, str):
        raise TypeError("Invalid object type")
    core = text.strip(_SPACE)
    match = _NUMBER.fullmatch(core)
    if match is not None:
        maybe_int = match.group(2) is None and match.group(3) is None and match.group(4) is None
        try:
            return float(core), maybe_int
        except (ValueError, OverflowError):
            pass
    if text.lower() in _INFINITY:
        return (-math.inf if text.startswith("-") else math.inf), False
    raise ValueError(f'Unable to parse string "{text}"')


def _numeric(values: list[Any], na: set[Any], masked: bool = False) -> _Values:
    """pandas' `maybe_convert_numeric`, which raises for anything not a number.

    With `masked`, as for a `dtype_backend`, a gap among integers or flags is
    kept as a gap in `mask` rather than turning the column into floats.
    """
    if not values:
        return _Values("int64", [], [] if masked else None)
    if _is_integer(values[0]):
        try:
            whole = []
            for value in values:
                if isinstance(value, (str, bytes)) or value is None:
                    raise TypeError
                number = int(value)
                if not _INT64[0] <= number <= _INT64[1] or number != value:
                    raise TypeError
                whole.append(number)
            return _Values("int64", whole, [False] * len(whole) if masked else None)
        except (TypeError, ValueError, OverflowError):
            pass
    seen_int = seen_bool = seen_null = seen_uint = seen_sint = False
    seen_overflow = seen_float = False
    floats: list[Any] = []
    ints: list[Any] = []
    mask: list[bool] = []
    allow_null_in_int = masked

    def saw_int(number: int) -> None:
        nonlocal seen_int, seen_sint, seen_uint
        seen_int = True
        seen_sint = seen_sint or _INT64[0] <= number < 0
        seen_uint = seen_uint or _INT64[1] < number <= _UINT64

    def saw_gap() -> None:
        nonlocal seen_null, seen_float
        seen_null = True
        if not allow_null_in_int:
            seen_float = True

    for at, value in enumerate(values):
        allow_null_in_int = masked and not seen_float
        gap = False
        if getattr(value, "__hash__", None) is not None and _in(value, na):
            saw_gap()
            gap = True
            floats.append(_NAN)
            ints.append(None)
        elif _is_float(value):
            number = float(value)
            if number != number:
                gap = True
                seen_null = True
                if not allow_null_in_int:
                    seen_float = True
            else:
                seen_float = True
            floats.append(number)
            ints.append(None)
        elif _is_integer(value):
            number = int(value)
            floats.append(float(number))
            saw_int(number)
            ints.append(number)
            if number > _UINT64 or number < _INT64[0] or (seen_sint and seen_uint):
                seen_overflow = True
        elif isinstance(value, bool):
            floats.append(float(value))
            ints.append(int(value))
            seen_bool = True
        elif value is None:
            saw_gap()
            gap = True
            floats.append(_NAN)
            ints.append(None)
        elif hasattr(value, "__len__") and len(value) == 0:
            raise InvalidArgumentError("Empty string encountered")
        elif type(value).__name__ == "Decimal":
            floats.append(float(value))
            ints.append(None)
            seen_float = True
        else:
            try:
                number, maybe_int = _floatify(value)
            except (TypeError, ValueError) as err:
                raise type(err)(f"{err} at position {at}") from None
            if _in(number, na):
                seen_null = seen_float = True
                gap = True
                floats.append(_NAN)
            else:
                if number != number:
                    seen_null = True
                    gap = True
                floats.append(number)
            ints.append(None)
            if maybe_int:
                whole = int(value)
                if _in(whole, na):
                    seen_null = seen_float = True
                    gap = True
                else:
                    saw_int(whole)
                    if whole < _INT64[0] or whole > _UINT64:
                        seen_overflow = True
                    else:
                        ints[-1] = whole
                seen_float = seen_float or (seen_uint and seen_sint)
            else:
                seen_float = True
        mask.append(gap)
    kept = mask if masked else None
    if seen_uint and (seen_null or seen_sint):
        return _Values("object", list(values))
    if allow_null_in_int and seen_null and not seen_int and not seen_bool:
        seen_float = True
    if seen_float:
        return _Values("float64", floats, kept)
    if seen_int:
        if seen_overflow:
            return _Values("object", ints, kept)
        return _Values("uint64" if seen_uint else "int64", ints, kept)
    if seen_bool:
        return _Values(
            "bool", [bool(v) for v in ints], mask if allow_null_in_int and masked else None
        )
    return _Values("int64", ints, kept)


def _sanitized(values: list[Any], na: set[Any]) -> tuple[list[Any], int]:
    """pandas' `sanitize_objects`: missing values as NaN, equal values as the first seen."""
    count = 0
    memo: dict[Any, Any] = {}
    out = []
    for value in values:
        try:
            hash(value)
        except TypeError:
            out.append(value)
            continue
        if _in(value, na):
            out.append(_NAN)
            count += 1
        elif value in memo:
            out.append(memo[value])
        else:
            memo[value] = value
            out.append(value)
    return out, count


def _as_bool(values: list[Any], truths: set[Any], falsities: set[Any]) -> _Values | None:
    """pandas' `maybe_convert_bool`, or None when a value is not a flag."""
    out = []
    gap = False
    for value in values:
        if isinstance(value, bool):
            out.append(value)
        elif _in(value, truths):
            out.append(True)
        elif _in(value, falsities):
            out.append(False)
        elif _missing(value):
            out.append(_NAN)
            gap = True
        else:
            return None
    return _Values("object" if gap else "bool", out)


def _converted_objects(values: list[Any]) -> _Values:
    """pandas' `maybe_convert_objects` on what a converter answered, numbers only."""
    present = [v for v in values if not _missing(v)]
    nulls = len(present) != len(values)
    if values and all(isinstance(v, bool) for v in values):
        return _Values("bool", list(values))
    if present and all(_is_integer(v) for v in present):
        if not nulls and all(_INT64[0] <= v <= _INT64[1] for v in present):
            return _Values("int64", [int(v) for v in values])
        if not nulls and all(0 <= v <= _UINT64 for v in present):
            return _Values("uint64", [int(v) for v in values])
        if nulls and all(_INT64[0] <= v <= _UINT64 for v in present):
            return _Values("float64", [_NAN if _missing(v) else float(v) for v in values])
        return _Values("object", list(values))
    if (
        present
        and all((_is_integer(v) or _is_float(v)) for v in present)
        and all(v is None or _is_float(v) or _is_integer(v) for v in values)
    ):
        return _Values("float64", [_NAN if _missing(v) else float(v) for v in values])
    return _Values("object", list(values))


class _TextParser:
    """pandas' `TextParser` with the Python engine, on a list of rows of values."""

    def __init__(self, data: list[list[Any]], options: dict[str, Any]) -> None:
        get = options.get
        self.nrows = _validate_integer("nrows", get("nrows"))
        skipfooter = get("skipfooter", 0)
        if skipfooter and self.nrows:
            raise InvalidArgumentError("'skipfooter' not supported with 'nrows'")
        header = get("header", 0)
        if header == "infer":
            header = 0 if get("names") is None else None
        _validate_header(header)
        index_col = get("index_col")
        if index_col is True:
            raise InvalidArgumentError("The value of index_col couldn't be 'True'")
        if (
            index_col is not None
            and index_col is not False
            and not isinstance(index_col, (list, tuple))
        ):
            index_col = [index_col]
        names = get("names")
        names = list(names) if names is not None else None
        converters = get("converters")
        if converters is not None:
            if not isinstance(converters, dict):
                raise TypeError(
                    "Type converters must be a dict or subclass, input was a"
                    f" {type(converters).__name__}"
                )
        else:
            converters = {}
        self.keep_default_na = get("keep_default_na", True)
        self.na_values, self.na_fvalues = _clean_na(get("na_values"), self.keep_default_na)
        skiprows = get("skiprows")
        if _is_integer(skiprows):
            skiprows = range(skiprows)
        if skiprows is None:
            skiprows = set()
        elif not callable(skiprows):
            skiprows = set(skiprows)
        self.skiprows = skiprows
        self.skipfunc = skiprows if callable(skiprows) else (lambda x: x in skiprows)

        # ParserBase
        self._implicit_index = False
        self.names = names
        self.orig_names: list[Any] | None = None
        self.index_col = index_col
        self.unnamed_cols: set[Any] = set()
        self.index_names: list[Any] | None = None
        self.col_names: list[Any] | None = None
        parse_dates = get("parse_dates", False)
        if parse_dates is None or isinstance(parse_dates, bool):
            parse_dates = bool(parse_dates)
        elif not isinstance(parse_dates, list):
            raise TypeError("Only booleans and lists are accepted for the 'parse_dates' parameter")
        self.parse_dates = parse_dates
        self.date_format = get("date_format")
        self.na_filter = get("na_filter", True)
        dtype = get("dtype")
        self.dtype = dict(dtype) if isinstance(dtype, dict) else dtype
        self.converters = converters
        self.dtype_backend = get("dtype_backend", NO_DEFAULT)
        self.true_values = get("true_values")
        self.false_values = get("false_values")
        self.header = header
        usecols = get("usecols")
        if _is_list_like(header, allow_sets=False):
            if usecols:
                raise InvalidArgumentError(
                    "cannot specify usecols when specifying a multi-index header"
                )
            if names:
                raise InvalidArgumentError(
                    "cannot specify names when specifying a multi-index header"
                )
            if self.index_col is not None:
                if _is_integer(self.index_col):
                    self.index_col = [self.index_col]
                elif not (
                    _is_list_like(self.index_col, allow_sets=False)
                    and all(map(_is_integer, self.index_col))
                ):
                    raise InvalidArgumentError(
                        "index_col must only contain integers of column positions when"
                        " specifying a multi-index header"
                    )
                else:
                    self.index_col = list(self.index_col)
        self.usecols, self.usecols_dtype = self._validate_usecols(usecols)

        # PythonParser
        self.data = data
        self.buf: list[list[Any]] = []
        self.pos = 0
        self.line_pos = 0
        if not _is_integer(skipfooter):
            raise InvalidArgumentError("skipfooter must be an integer")
        if skipfooter < 0:
            raise InvalidArgumentError("skipfooter cannot be negative")
        self.skipfooter = skipfooter
        self.has_index_names = get("has_index_names", False)
        self.thousands = get("thousands")
        self.decimal = get("decimal", ".")
        self.comment = get("comment")
        self._col_indices: list[int] | None = None
        self._header_line_cache: Any = NO_DEFAULT
        columns, self.num_original_columns, self.unnamed_cols = self._infer_columns()
        self.columns, self.index_names, self.col_names = self._multi_indexer_columns(
            columns, self.index_names
        )
        self.orig_names = list(self.columns)
        index_names, self.orig_names, self.columns = self._index_name()
        if self.index_names is None:
            self.index_names = index_names
        if self._col_indices is None:
            self._col_indices = list(range(len(self.columns)))
        self._no_thousands_columns = self._no_thousand_columns()
        if len(self.decimal) != 1:
            raise InvalidArgumentError("Only length-1 decimal markers supported")
        decimal = re.escape(self.decimal)
        if self.thousands is None:
            pattern = f"^[\\-\\+]?[0-9]*({decimal}[0-9]*)?([0-9]?(E|e)\\-?[0-9]+)?$"
        else:
            thousands = re.escape(self.thousands)
            pattern = (
                f"^[\\-\\+]?([0-9]+{thousands}|[0-9])*({decimal}[0-9]*)?([0-9]?(E|e)\\-?[0-9]+)?$"
            )
        self.num = re.compile(pattern)

    @staticmethod
    def _validate_usecols(usecols: Any) -> tuple[Any, str | None]:
        if usecols is None:
            return None, None
        if callable(usecols):
            return usecols, None
        if not _is_list_like(usecols):
            raise InvalidArgumentError(_USECOLS)
        listed = list(usecols)
        if not listed:
            kind = "empty"
        elif all(_is_integer(u) for u in listed):
            kind = "integer"
        elif all(isinstance(u, str) for u in listed):
            kind = "string"
        else:
            raise InvalidArgumentError(_USECOLS)
        return set(listed), kind

    # Reading the rows.

    def _check_comments(self, lines: list[list[Any]]) -> list[list[Any]]:
        if self.comment is None:
            return lines
        out = []
        for line in lines:
            kept = []
            for x in line:
                if not isinstance(x, str) or self.comment not in x or _in(x, self.na_values):
                    kept.append(x)
                else:
                    x = x[: x.find(self.comment)]
                    if len(x) > 0:
                        kept.append(x)
                    break
            out.append(kept)
        return out

    @staticmethod
    def _is_line_empty(line: Sequence[Any]) -> bool:
        return not line or all(not x for x in line)

    def _next_line(self) -> list[Any]:
        while self.skipfunc(self.pos):
            if self.pos >= len(self.data):
                break
            self.pos += 1
        while True:
            if self.pos >= len(self.data):
                raise StopIteration
            line = self._check_comments([self.data[self.pos]])[0]
            self.pos += 1
            if self._is_line_empty(self.data[self.pos - 1]) or line:
                break
        self.line_pos += 1
        self.buf.append(line)
        return line

    def _buffered_line(self) -> list[Any]:
        return self.buf[0] if self.buf else self._next_line()

    def _header_line(self) -> Any:
        if self._header_line_cache is NO_DEFAULT:
            if self.header is not None:
                self._header_line_cache = None
            else:
                try:
                    line = self._buffered_line()
                except StopIteration:
                    if not self.names:
                        raise EmptyDataError("No columns to parse from file") from None
                    line = self.names[:]
                self._header_line_cache = line
        return self._header_line_cache

    def _have_mi_columns(self) -> bool:
        return isinstance(self.header, (list, tuple)) and len(self.header) > 1

    def _infer_columns(self) -> tuple[list[list[Any]], int, set[Any]]:
        names = self.names
        num_original_columns = 0
        clear_buffer = True
        unnamed_cols: set[Any] = set()
        if self.header is not None:
            header = self.header
            have_mi_columns = self._have_mi_columns()
            if isinstance(header, (list, tuple)):
                if have_mi_columns:
                    header = [*list(header), header[-1] + 1]
            else:
                header = [header]
            columns: list[list[Any]] = []
            for level, hr in enumerate(header):
                try:
                    line = self._buffered_line()
                    while self.line_pos <= hr:
                        line = self._next_line()
                except StopIteration:
                    if 0 < self.line_pos <= hr and (not have_mi_columns or hr != header[-1]):
                        joined = list(map(str, header[:-1] if have_mi_columns else header))
                        message = f"[{','.join(joined)}], len of {len(joined)}, "
                        raise InvalidArgumentError(
                            f"Passed header={message}but only {self.line_pos} lines in file"
                        ) from None
                    if have_mi_columns and hr > 0:
                        if clear_buffer:
                            self.buf.clear()
                        columns.append([None] * len(columns[-1]))
                        return columns, num_original_columns, unnamed_cols
                    if not self.names:
                        raise EmptyDataError("No columns to parse from file") from None
                    line = self.names[:]
                this_columns: list[Any] = []
                this_unnamed = []
                for i, c in enumerate(line):
                    if isinstance(c, str) and c == "":
                        if have_mi_columns:
                            this_columns.append(f"Unnamed: {i}_level_{level}")
                        else:
                            this_columns.append(f"Unnamed: {i}")
                        this_unnamed.append(i)
                    else:
                        this_columns.append(c)
                if not have_mi_columns:
                    counts: defaultdict[Any, int] = defaultdict(int)
                    order = [
                        i for i in range(len(this_columns)) if i not in this_unnamed
                    ] + this_unnamed
                    for i in order:
                        col = this_columns[i]
                        old = col
                        count = counts[col]
                        if count > 0:
                            while count > 0:
                                counts[old] = count + 1
                                col = f"{old}.{count}"
                                if col in this_columns:
                                    count += 1
                                else:
                                    count = counts[col]
                            if (
                                isinstance(self.dtype, dict)
                                and self.dtype.get(old) is not None
                                and self.dtype.get(col) is None
                            ):
                                self.dtype[col] = self.dtype[old]
                        this_columns[i] = col
                        counts[col] = count + 1
                elif hr == header[-1]:
                    width = len(this_columns)
                    ic = len(self.index_col) if self.index_col is not None else 0
                    unnamed_count = len(this_unnamed)
                    if (width != unnamed_count and width - ic > unnamed_count) or ic == 0:
                        clear_buffer = False
                        this_columns = [None] * width
                        self.buf = [self.buf[-1]]
                columns.append(this_columns)
                unnamed_cols.update({this_columns[i] for i in this_unnamed})
                if len(columns) == 1:
                    num_original_columns = len(this_columns)
            if clear_buffer:
                self.buf.clear()
            if names is not None:
                try:
                    first_line: list[Any] | None = self._next_line()
                except StopIteration:
                    first_line = None
                first_width = 0 if first_line is None else len(first_line)
                if len(names) > len(columns[0]) and len(names) > first_width:
                    raise InvalidArgumentError(
                        "Number of passed names did not match number of header fields in the file"
                    )
                if len(columns) > 1:
                    raise TypeError("Cannot pass names with multi-index columns")
                if self.usecols is not None:
                    self._handle_usecols(columns, names, num_original_columns)
                else:
                    num_original_columns = len(names)
                if self._col_indices is not None and len(names) != len(self._col_indices):
                    columns = [[names[i] for i in sorted(self._col_indices)]]
                else:
                    columns = [names]
            else:
                columns = self._handle_usecols(columns, columns[0], num_original_columns)
        else:
            width = len(self._header_line())
            num_original_columns = width
            if not names:
                columns = [list(range(width))]
                columns = self._handle_usecols(columns, columns[0], width)
            elif self.usecols is None or len(names) >= width:
                columns = self._handle_usecols([names], names, width)
                num_original_columns = len(names)
            elif not callable(self.usecols) and len(names) != len(self.usecols):
                raise InvalidArgumentError(
                    "Number of passed names did not match number of header fields in the file"
                )
            else:
                columns = [names]
                self._handle_usecols(columns, columns[0], width)
        return columns, num_original_columns, unnamed_cols

    def _handle_usecols(
        self, columns: list[list[Any]], key: list[Any], num_original_columns: int
    ) -> list[list[Any]]:
        if self.usecols is None:
            return columns
        indices: Any
        if callable(self.usecols):
            indices = {i for i, name in enumerate(key) if self.usecols(name)}
        elif any(isinstance(u, str) for u in self.usecols):
            if len(columns) > 1:
                raise InvalidArgumentError("If using multiple headers, usecols must be integers.")
            indices = []
            for col in self.usecols:
                if isinstance(col, str):
                    try:
                        indices.append(key.index(col))
                    except ValueError:
                        missing = [c for c in self.usecols if c not in key]
                        if missing:
                            raise InvalidArgumentError(
                                "Usecols do not match columns, columns expected but not found:"
                                f" {missing}"
                            ) from None
                else:
                    indices.append(col)
        else:
            beyond = [col for col in self.usecols if col >= num_original_columns]
            if beyond:
                raise ParserError(
                    "Defining usecols with out-of-bounds indices is not allowed."
                    f" {beyond} are out-of-bounds."
                )
            indices = self.usecols
        columns = [[n for i, n in enumerate(column) if i in indices] for column in columns]
        self._col_indices = sorted(indices)
        return columns

    def _clean_index_names(self, columns: Any, index_col: Any) -> tuple[Any, list, Any]:
        if index_col is None or index_col is False:
            return None, columns, index_col
        columns = list(columns)
        if not columns:
            return [None] * len(index_col), columns, index_col
        copied = list(columns)
        index_names: list[Any] = []
        index_col = list(index_col)
        for i, c in enumerate(index_col):
            if isinstance(c, str):
                index_names.append(c)
                for j, name in enumerate(copied):
                    if name == c:
                        index_col[i] = j
                        columns.remove(name)
                        break
            else:
                name = copied[c]
                columns.remove(name)
                index_names.append(name)
        for i, name in enumerate(index_names):
            if isinstance(name, str) and name in self.unnamed_cols:
                index_names[i] = None
        return index_names, columns, index_col

    def _multi_indexer_columns(self, header: list[list[Any]], index_names: Any) -> tuple:
        if len(header) < 2:
            return header[0], index_names, None
        ic = self.index_col
        if ic is None:
            ic = []
        if not isinstance(ic, (list, tuple)):
            ic = [ic]
        skipped = set(ic)
        index_names = header.pop(-1)
        index_names, _, _ = self._clean_index_names(index_names, self.index_col)
        width = len(header[0])
        if not all(len(row) == width for row in header[1:]):
            raise ParserError("Header rows must have an equal number of columns.")
        columns = list(
            zip(
                *(tuple(r[i] for i in range(width) if i not in skipped) for r in header),
                strict=True,
            )
        )
        names = columns.copy()
        for single in sorted(ic):
            names.insert(single, single)
        if ic:
            col_names = [
                r[ic[0]] if r[ic[0]] is not None and not _in(r[ic[0]], self.unnamed_cols) else None
                for r in header
            ]
        else:
            col_names = [None] * len(header)
        return names, index_names, col_names

    def _index_name(self) -> tuple[Any, list[Any], list[Any]]:
        columns = list(self.orig_names or [])
        orig_names = list(columns)
        line: list[Any] | None
        if self._header_line() is not None:
            line = self._header_line()
        else:
            try:
                line = self._next_line()
            except StopIteration:
                line = None
        try:
            next_line: list[Any] | None = self._next_line()
        except StopIteration:
            next_line = None
        implicit_first_cols = 0
        if line is not None:
            index_col = self.index_col
            if index_col is not False:
                implicit_first_cols = len(line) - self.num_original_columns
            if (
                next_line is not None
                and self.header is not None
                and index_col is not False
                and len(next_line) == len(line) + self.num_original_columns
            ):
                self.index_col = list(range(len(line)))
                self.buf = self.buf[1:]
                for c in reversed(line):
                    columns.insert(0, c)
                orig_names = list(columns)
                self.num_original_columns = len(columns)
                return line, orig_names, columns
        if implicit_first_cols > 0:
            self._implicit_index = True
            if self.index_col is None:
                self.index_col = list(range(implicit_first_cols))
            index_name = None
        else:
            index_name, _, self.index_col = self._clean_index_names(columns, self.index_col)
        return index_name, orig_names, columns

    def _no_thousand_columns(self) -> set[int]:
        skipped: set[int] = set()
        indices = self._col_indices or []
        if self.columns and self.parse_dates:
            skipped = self._noconvert_columns(indices, self.columns)
        if self.columns and self.dtype:
            for i, col in zip(indices, self.columns, strict=True):
                if not isinstance(self.dtype, dict) and not _numeric_dtype(self.dtype):
                    skipped.add(i)
                if isinstance(self.dtype, dict) and col in self.dtype:
                    wanted = self.dtype[col]
                    if not _numeric_dtype(wanted) or _dtype_name(wanted) == "bool":
                        skipped.add(i)
        return skipped

    def _noconvert_columns(self, indices: list[int], names: Sequence[Any]) -> set[int]:
        if self.usecols_dtype == "integer":
            usecols: list[int] | None = sorted(self.usecols)
        elif callable(self.usecols) or self.usecols_dtype not in ("empty", None):
            usecols = indices
        else:
            usecols = None

        def position(x: Any) -> int:
            if usecols is not None and _is_integer(x):
                x = usecols[x]
            if not _is_integer(x):
                x = indices[list(names).index(x)]
            return x

        out = set()
        if isinstance(self.parse_dates, list):
            _dates_present(self.parse_dates, names)
            for value in self.parse_dates:
                out.add(position(value))
        elif self.parse_dates:
            if isinstance(self.index_col, list):
                for k in self.index_col:
                    out.add(position(k))
            elif self.index_col is not None:
                out.add(position(self.index_col))
        return out

    def _get_lines(self, rows: int | None) -> list[list[Any]]:
        lines = self.buf
        new_rows = None
        if rows is not None:
            if len(self.buf) >= rows:
                new_rows, self.buf = self.buf[:rows], self.buf[rows:]
            else:
                rows -= len(self.buf)
        if new_rows is None:
            if self.pos > len(self.data):
                raise StopIteration
            if rows is None:
                new_rows = self.data[self.pos :]
                new_pos = len(self.data)
            else:
                new_rows = self.data[self.pos : self.pos + rows]
                new_pos = self.pos + rows
            if self.skiprows:
                new_rows = [
                    row for i, row in enumerate(new_rows) if not self.skipfunc(i + self.pos)
                ]
            lines.extend(new_rows)
            self.pos = new_pos
            self.buf = []
        else:
            lines = new_rows
        if self.skipfooter:
            lines = lines[: -self.skipfooter]
        lines = self._check_comments(lines)
        if self.thousands is not None:
            lines = self._replace_in_numbers(lines, self.thousands, "")
        if self.decimal != ".":
            lines = self._replace_in_numbers(lines, self.decimal, ".")
        return lines

    def _replace_in_numbers(
        self, lines: list[list[Any]], search: str, replace: str
    ) -> list[list[Any]]:
        out = []
        for line in lines:
            kept = []
            for i, x in enumerate(line):
                if (
                    not isinstance(x, str)
                    or search not in x
                    or i in self._no_thousands_columns
                    or not self.num.search(x.strip())
                ):
                    kept.append(x)
                else:
                    kept.append(x.replace(search, replace))
            out.append(kept)
        return out

    def _rows_to_cols(self, content: list[list[Any]]) -> list[list[Any]]:
        width = self.num_original_columns
        if self._implicit_index:
            width += len(self.index_col)
        longest = max(len(row) for row in content)
        if longest > width and self.index_col is not False and self.usecols is None:
            footers = self.skipfooter or 0
            total = len(content)
            for i, row in enumerate(content):
                if len(row) > width:
                    row_num = self.pos - (total - i + footers)
                    raise ParserError(
                        f"Expected {width} fields in line {row_num + 1}, saw {len(row)}"
                    )
        width = max(width, longest)
        columns = [[row[i] if i < len(row) else None for row in content] for i in range(width)]
        if self.usecols:
            indices = self._col_indices or []
            if self._implicit_index:
                lead = len(self.index_col)
                columns = [a for i, a in enumerate(columns) if i < lead or i - lead in indices]
            else:
                columns = [a for i, a in enumerate(columns) if i in indices]
        return columns

    # Typing the columns.

    def _infer_types(
        self, values: Any, na: set[Any], no_dtype: bool, try_num_bool: bool = True
    ) -> tuple[_Values, int]:
        if isinstance(values, _Values) and values.kind != "object":
            wanted = {v for v in na if not isinstance(v, str)}
            hits = [_in(v, wanted) for v in values.values]
            count = sum(hits)
            if count:
                kind = "float64" if values.kind.startswith(("int", "uint")) else values.kind
                if kind == "bool":
                    kind = "object"
                values = _Values(
                    kind,
                    [
                        _NAN if hit else (float(v) if kind == "float64" else v)
                        for v, hit in zip(values.values, hits, strict=True)
                    ],
                )
            return values, count
        raw = values.values if isinstance(values, _Values) else list(values)
        masked = no_dtype and self.dtype_backend is not NO_DEFAULT
        result: _Values
        count = 0
        if try_num_bool:
            try:
                result = _numeric(raw, na, masked)
            except (ValueError, TypeError):
                raw, count = _sanitized(raw, na)
                result = _Values("object", raw)
            else:
                if masked:
                    gaps = result.mask or [False] * len(result.values)
                    count = sum(gaps)
                    if gaps and all(gaps):
                        result = _Values("Int64", [1] * len(gaps), gaps)
                    elif result.kind in ("int64", "uint64", "float64", "bool"):
                        kind = {"int64": "Int64", "uint64": "UInt64", "float64": "Float64"}
                        result = _Values(kind.get(result.kind, "boolean"), result.values, gaps)
                else:
                    count = sum(1 for v in result.values if _missing(v))
        else:
            raw, count = _sanitized(raw, na)
            result = _Values("object", raw)
        if (
            result.kind == "object"
            and try_num_bool
            and (not result.values or not isinstance(result.values[0], int))
        ):
            truths = set(_TRUE) | set(self.true_values or ())
            falsities = set(_FALSE) | set(self.false_values or ())
            flagged = _as_bool(result.values, truths, falsities)
            if flagged is not None and masked:
                gaps = [_missing(v) for v in flagged.values]
                result = _Values(
                    "boolean",
                    [
                        bool(v) if not g else False
                        for v, g in zip(flagged.values, gaps, strict=True)
                    ],
                    gaps,
                )
            elif flagged is not None:
                result = flagged
            elif masked:
                present = [v for v in result.values if not _missing(v)]
                if not all(isinstance(v, dt.datetime) for v in present):
                    texts = [None if _missing(v) else str(v) for v in raw]
                    result = _Values("string", texts, [v is None for v in texts])
        return result, count

    def _clean_mapping(self, mapping: Any) -> Any:
        if not isinstance(mapping, dict):
            return mapping
        clean = {}
        for col, value in mapping.items():
            if isinstance(col, int) and col not in (self.orig_names or []):
                col = self.orig_names[col]  # type: ignore[index]
            clean[col] = value
        if isinstance(mapping, defaultdict):
            remaining = set(self.orig_names or []) - set(clean)
            clean.update({col: mapping[col] for col in remaining})
        return clean

    def _convert(self, data: dict[Any, list[Any]]) -> dict[Any, Any]:
        converters = self._clean_mapping(self.converters)
        dtypes = self._clean_mapping(self.dtype)
        if isinstance(self.na_values, dict):
            na_values: Any = {}
            na_fvalues: Any = {}
            for col in self.na_values:
                if col is not None:
                    value = self.na_values[col]
                    fvalue = self.na_fvalues[col]
                    if isinstance(col, int) and col not in (self.orig_names or []):
                        col = self.orig_names[col]  # type: ignore[index]
                    na_values[col] = value
                    na_fvalues[col] = fvalue
        else:
            na_values, na_fvalues = self.na_values, self.na_fvalues
        dated = _dates_present(self.parse_dates, self.columns)
        result = {}
        for c, values in data.items():
            convert = None if converters is None else converters.get(c)
            cast = dtypes.get(c) if isinstance(dtypes, dict) else dtypes
            if self.na_filter:
                col_na, col_fna = _na_for(c, na_values, na_fvalues, self.keep_default_na)
            else:
                col_na, col_fna = set(), set()
            both = set(col_na) | set(col_fna)
            if c in dated:
                result[c] = _Values("object", [_NAN if _in(v, both) else v for v in values])
                continue
            if convert is not None:
                if cast is not None:
                    warnings.warn(
                        f"Both a converter and dtype were specified for column {c} - only the"
                        " converter will be used.",
                        ParserWarning,
                        stacklevel=6,
                    )
                try:
                    mapped = [convert(v) for v in values]
                except ValueError:
                    listed = set(na_values) if not isinstance(na_values, dict) else set()
                    mapped = [v if _in(v, listed) else convert(v) for v in values]
                column, _ = self._infer_types(
                    _converted_objects(mapped), both, cast is None, try_num_bool=False
                )
            else:
                textual = cast is not None and (_extension_dtype(cast) or _string_dtype(cast))
                column, count = self._infer_types(values, both, cast is None, not textual)
                if cast is not None:
                    column = _Cast(column, cast, c, count, self)
            result[c] = column
        return result

    def _dates(self, names: Sequence[Any], data: dict[Any, Any]) -> dict[Any, Any]:
        if not isinstance(self.parse_dates, list):
            return data
        for spec in self.parse_dates:
            if isinstance(spec, int) and spec not in data:
                spec = names[spec]
            if (isinstance(self.index_col, list) and spec in self.index_col) or (
                isinstance(self.index_names, list) and spec in self.index_names
            ):
                continue
            data[spec] = _date_converted(data[spec], spec, self.date_format)
        return data

    def _agg_index(self, index: list[list[Any]]) -> Any:
        from ._frame import Index
        from ._multi import MultiIndex

        converters = self._clean_mapping(self.converters)
        dtypes = self._clean_mapping(self.dtype)
        names: Iterable[Any] = self.index_names if self.index_names is not None else None  # type: ignore[assignment]
        built = []
        for i, values in enumerate(index):
            name = names[i] if names is not None else None  # type: ignore[index]
            column: Any = _Values("object", list(values))
            if self._should_parse_dates(i):
                column = _date_converted(column, name, self.date_format)
            if self.na_filter:
                col_na, col_fna = self.na_values, self.na_fvalues
            else:
                col_na, col_fna = set(), set()
            if isinstance(self.na_values, dict):
                if name is not None:
                    col_na, col_fna = _na_for(
                        name, self.na_values, self.na_fvalues, self.keep_default_na
                    )
                else:
                    col_na, col_fna = set(), set()
            cast = None
            converted = False
            if self.index_names is not None:
                if isinstance(dtypes, dict):
                    cast = dtypes.get(name)
                if isinstance(converters, dict):
                    converted = converters.get(name) is not None
            try_num_bool = not ((cast and _string_dtype(cast)) or converted)
            if isinstance(column, _Values):
                column, _ = self._infer_types(
                    column, set(col_na) | set(col_fna), cast is None, try_num_bool
                )
            if cast is not None and isinstance(column, _Values) and column.kind == "object":
                # pandas casts the index's objects before they are read as moments.
                series = _Cast(column, cast, name, 0, self)
                series = _series(series, backend=NO_DEFAULT)
            else:
                series = _series(column, backend=NO_DEFAULT)
                if cast is not None:
                    series = series.astype(cast)
            built.append(Index(series, name=name))
        if len(built) == 1:
            return built[0]
        return MultiIndex.from_arrays(built)

    def _should_parse_dates(self, i: int) -> bool:
        if isinstance(self.parse_dates, bool):
            return self.parse_dates
        name = self.index_names[i] if self.index_names is not None else None
        j = i if self.index_col is None else self.index_col[i]
        return j in self.parse_dates or (name is not None and name in self.parse_dates)

    def _make_index(self, alldata: list[Any], columns: list[Any], indexnamerow: Any) -> tuple:
        index = None
        if isinstance(self.index_col, list) and len(self.index_col):
            removed = []
            indexes = []
            for idx in self.index_col:
                if isinstance(idx, str):
                    raise InvalidArgumentError(f"Index {idx} invalid")
                removed.append(idx)
                indexes.append(alldata[idx])
            for i in sorted(removed, reverse=True):
                alldata.pop(i)
                if not self._implicit_index:
                    columns.pop(i)
            index = self._agg_index(indexes)
            if indexnamerow:
                offset = len(indexnamerow) - len(columns)
                index = index.set_names(indexnamerow[:offset])
        return index, self._multi_columns(columns)

    def _multi_columns(self, columns: list[Any]) -> Any:
        if _potential_multi(columns):
            from ._multi import MultiIndex

            return MultiIndex.from_tuples(columns, names=self.col_names)
        return columns

    def _empty(self, columns: list[Any]) -> Any:
        from ._frame import DataFrame, Index
        from ._multi import MultiIndex
        from ._range_index import RangeIndex

        columns = list(columns)
        dtype = self.dtype
        if not isinstance(dtype, dict):
            by_name: defaultdict[Any, Any] = defaultdict(lambda: dtype)
        else:
            by_name = defaultdict(
                lambda: None,
                {columns[k] if _is_integer(k) else k: v for k, v in dtype.items()},
            )
        if self.index_col is None or self.index_col is False or self.index_names is None:
            index: Any = RangeIndex(0)
        else:
            levels = [Index(_empty_series(by_name[name]), name=name) for name in self.index_names]
            index = levels[0] if len(levels) == 1 else MultiIndex.from_arrays(levels)
            self.index_col.sort()
            for i, n in enumerate(self.index_col):
                columns.pop(n - i)
        frame = DataFrame(
            {at: _empty_series(by_name[name]) for at, name in enumerate(columns)},
            index=index,
        )
        frame = frame.set_axis(self._multi_columns(columns), axis=1)
        return frame

    def read(self, rows: int | None) -> Any:
        from ._frame import DataFrame
        from ._range_index import RangeIndex

        try:
            content = self._get_lines(rows)
        except StopIteration:
            content = []
        columns = list(self.orig_names or [])
        if not content:
            names = _dedup(
                self.orig_names or [], _potential_multi(self.orig_names or [], self.index_col)
            )
            return self._empty(names)
        indexnamerow = None
        if self.has_index_names and sum(int(_blank_cell(v)) for v in content[0]) == len(columns):
            indexnamerow = content[0]
            content = content[1:]
        alldata = self._rows_to_cols(content)
        names = _dedup(
            self.orig_names or [], _potential_multi(self.orig_names or [], self.index_col)
        )
        offset = len(self.index_col) if self._implicit_index else 0
        self._check_data_length(names, alldata)
        data = {name: alldata[i + offset] for i, name in enumerate(names) if i < len(alldata)}
        converted = self._convert(data)
        converted = self._dates(names, converted)
        index, result_columns = self._make_index(alldata, names, indexnamerow)
        length = len(content)
        backend = self.dtype_backend
        wanted = self.dtype
        built = {}
        for at, name in enumerate(list(result_columns)):
            column = converted[name]
            as_object = (
                isinstance(wanted, dict) and _dtype_name(wanted.get(name)) in ("object", "O")
            ) or (not isinstance(wanted, dict) and _dtype_name(wanted) in ("object", "O"))
            built[at] = _series(column, backend=backend, as_object=as_object)
        if index is None:
            index = RangeIndex(0, length)
        # The columns are built on their own positions, so the index is put on after.
        frame = DataFrame(built).set_axis(index, axis=0) if built else DataFrame(index=index)
        frame = frame.set_axis(_columns_index(result_columns, list(converted)), axis=1)
        # Labels firepanda cannot hold are refused on reading them, so they are read here.
        _ = frame.columns
        return frame

    def _check_data_length(self, columns: Sequence[Any], data: list[list[Any]]) -> None:
        if not self.index_col and len(columns) != len(data) and columns:
            last = data[-1]
            if len(columns) == len(data) - 1 and all(
                _missing(v) or (isinstance(v, str) and v == "") for v in last
            ):
                return
            warnings.warn(
                "Length of header or names does not match length of data. This leads to a loss"
                " of data with index_col=False.",
                ParserWarning,
                stacklevel=6,
            )


def _columns_index(columns: Any, keys: list[Any]) -> Any:
    from ._frame import Index

    if hasattr(columns, "nlevels"):
        return columns
    labels = list(columns)
    if labels and (
        all(isinstance(v, str) for v in labels)
        or all(_is_integer(v) for v in labels)
        or all(isinstance(v, float) for v in labels)
    ):
        return Index(labels)
    return Index(_objects_series(labels))


def _dates_present(parse_dates: Any, columns: Sequence[Any]) -> set[Any]:
    if not isinstance(parse_dates, list):
        return set()
    missing = set()
    found = set()
    for col in parse_dates:
        if isinstance(col, str):
            if col not in columns:
                missing.add(col)
            else:
                found.add(col)
        elif col in columns:
            found.add(col)
        else:
            found.add(columns[col])
    if missing:
        joined = ", ".join(sorted(missing))
        raise InvalidArgumentError(f"Missing column provided to 'parse_dates': '{joined}'")
    return found


def _date_converted(column: Any, name: Any, date_format: Any) -> Any:
    """pandas' `date_converter`: text read as moments, or left as text when it is not."""
    from ._frame import Series
    from ._pandas import to_datetime

    if not isinstance(column, _Values):
        printed = str(column.dtype)
        if printed.startswith(("datetime64", "timedelta64")):
            return column
        values = column.tolist()
    else:
        values = column.values
    fmt = date_format.get(name) if isinstance(date_format, dict) else date_format
    text = [v if _missing(v) else str(v) for v in values]
    try:
        return to_datetime(Series(text), format=fmt)
    except (ValueError, TypeError):
        return _Values("object", text)


def _dtype_name(dtype: Any) -> str:
    if dtype is None:
        return ""
    if dtype is object:
        return "object"
    if dtype is str:
        return "str"
    if isinstance(dtype, type):
        numpy_names = {int: "int64", float: "float64", complex: "complex128"}
        return numpy_names.get(dtype, dtype.__name__)
    return str(dtype)


def _numeric_dtype(dtype: Any) -> bool:
    name = _dtype_name(dtype)
    if dtype in (int, float, bool):
        return True
    return name.lower().startswith(("int", "uint", "float", "bool", "complex"))


def _string_dtype(dtype: Any) -> bool:
    return _dtype_name(dtype) in ("str", "string", "object", "O") or _dtype_name(dtype).startswith(
        ("string[", "<U", "|S")
    )


def _extension_dtype(dtype: Any) -> bool:
    name = _dtype_name(dtype)
    return (
        name.startswith(("Int", "UInt", "Float", "category", "string", "str", "Sparse"))
        or name in ("boolean", "category")
        or "[pyarrow]" in name
        or type(dtype).__name__.endswith("Dtype")
    )


def _Cast(column: _Values, cast: Any, name: Any, count: int, parser: _TextParser) -> Any:
    """pandas' `_cast_types`: a parsed column turned into the type `dtype` asked for."""
    from ._frame import Series

    wanted = _dtype_name(cast)
    if wanted in ("bool",) and count > 0 and not _extension_dtype(cast):
        raise InvalidArgumentError(f"Bool column has NA values in column {name}")
    if wanted == column.kind:
        return column
    if wanted in ("object", "O"):
        return column
    if wanted in ("str", "string") or wanted.startswith("string["):
        values = [v if _missing(v) else str(v) for v in column.values]
        return Series(values, dtype=cast)
    if wanted == "boolean":
        truths = set(_TRUE) | {"1", "1.0"} | set(parser.true_values or ())
        falsities = set(_FALSE) | {"0", "0.0"} | set(parser.false_values or ())
        flags: list[Any] = []
        for value in column.values:
            text = str(value)
            if _missing(value) or _in(text, parser.na_values):
                flags.append(None)
            elif text in truths:
                flags.append(True)
            elif text in falsities:
                flags.append(False)
            else:
                raise InvalidArgumentError(f"{text} cannot be cast to bool")
        return Series(flags, dtype="boolean")
    if wanted.startswith(("Int", "UInt", "Float")):
        numbers_ = _numeric([_NAN if _missing(v) else v for v in column.values], set())
        return _series(numbers_, NO_DEFAULT).astype(cast)
    if wanted.startswith("category") or type(cast).__name__ == "CategoricalDtype":
        return _series(column, NO_DEFAULT, as_object=True).astype(cast)
    try:
        if column.kind == "object" and wanted.startswith(("int", "uint", "float")):
            # numpy casts an object array with `int` or `float` on each value, so a value
            # neither takes raises numpy's TypeError, which pandas lets through.
            convert = float if wanted.startswith("float") else int
            return Series([convert(v) for v in column.values], dtype=wanted)
        return _series(column, NO_DEFAULT).astype(cast)
    except ValueError as err:
        raise InvalidArgumentError(f"Unable to convert column {name} to type {wanted}") from err


def _empty_series(dtype: Any) -> Any:
    """An empty column of a type, an object column when none is asked for."""
    from ._frame import Series

    if dtype is None or _dtype_name(dtype) in ("object", "O"):
        return _objects_series([])
    return Series([], dtype=dtype)


def _objects_series(values: list[Any]) -> Any:
    """An object column holding the values as they are."""
    from ._frame import Series

    return Series(_objects.cells(values), dtype="str")


def _series(column: Any, backend: Any, as_object: bool = False) -> Any:
    """A parsed column as a firepanda series, typed as pandas' frame constructor types it."""
    from ._frame import Series

    if not isinstance(column, _Values):
        return column
    kind = column.kind
    values = column.values
    if column.mask is not None and kind in ("Int64", "UInt64", "Float64", "boolean", "string"):
        series = Series(
            [None if gap else v for v, gap in zip(values, column.mask, strict=True)], dtype=kind
        )
    elif kind in ("int64", "uint64", "float64", "bool"):
        series = Series(values, dtype=kind)
    elif as_object:
        series = _objects_series(values)
    else:
        series = _inferred(values)
    if backend == "pyarrow":
        series = _arrow_backed(column)
    return series


_ARROW_TYPES = {
    "Int64": "int64",
    "UInt64": "uint64",
    "Float64": "float64",
    "boolean": "bool_",
    "int64": "int64",
    "uint64": "uint64",
    "float64": "float64",
    "bool": "bool_",
}


def _arrow_backed(column: _Values) -> Any:
    """A parsed column backed by Arrow, as pandas builds it for `dtype_backend="pyarrow"`."""
    import pyarrow as pa

    from ._arrowtyped import arrow_series

    gaps = column.mask or [False] * len(column.values)
    if gaps and all(gaps):
        return arrow_series(pa.array([None] * len(gaps)))
    values = [None if gap else v for v, gap in zip(column.values, gaps, strict=True)]
    name = _ARROW_TYPES.get(column.kind)
    if name is not None:
        return arrow_series(pa.array(values, type=getattr(pa, name)(), from_pandas=True))
    return arrow_series(pa.array(values, from_pandas=True))


def _inferred(values: list[Any]) -> Any:
    """pandas' frame constructor on an object column: text, moments or objects."""
    from ._frame import Series

    present = [v for v in values if not _missing(v)]
    if present and (
        all(isinstance(v, str) for v in present)
        or all(isinstance(v, dt.datetime) and v.tzinfo is None for v in present)
        or all(isinstance(v, dt.timedelta) for v in present)
    ):
        return Series([None if _missing(v) else v for v in values])
    return _objects_series(values)
