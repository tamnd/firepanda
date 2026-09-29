"""`ExcelWriter` and `to_excel`, which write a frame to a workbook the way pandas does.

pandas writes a workbook in two steps. `ExcelFormatter` turns the frame into
cells, each a row, a column, a value and maybe a merged span: the header rows,
the index names, the index and the body. Then a writer for the engine puts
those cells into openpyxl, xlsxwriter or odfpy and saves the book. This module
ports both steps, so the cells land where pandas puts them, carry pandas'
number formats, and fail with pandas' sentences. The engines are the same
packages pandas drives, so the file each one saves is the file pandas would
have saved.

The formatter is pandas' own. Given a frame every cell has no style, and given a
`Styler` each cell carries the Excel style pandas' `CSSToExcelConverter` makes of
the CSS the Styler computed for it, which each writer turns into its engine's own
fonts, fills and borders.
"""

from __future__ import annotations

import datetime as dt
import decimal
import functools
import itertools
import json
import math
import numbers
import os
import warnings
from collections import defaultdict
from collections.abc import Iterator
from pathlib import Path
from typing import Any, ClassVar

from . import _optional

_STORAGE = "storage_options passed with file object or non-fsspec file path"
_IF_SHEET_EXISTS = (
    "'{}' is not valid for if_sheet_exists. "
    "Valid options are 'error', 'new', 'replace' and 'overlay'."
)
_WRITERS: dict[str, type[ExcelWriter]] = {}


def _is_integer(value: Any) -> bool:
    return isinstance(value, numbers.Integral) and not _is_bool(value)


def _is_float(value: Any) -> bool:
    return isinstance(value, float) or (
        type(value).__module__ == "numpy"
        and isinstance(value, numbers.Real)
        and not isinstance(value, numbers.Integral)
    )


def _is_bool(value: Any) -> bool:
    return isinstance(value, bool) or type(value).__name__ == "bool_"


def _missing(value: Any) -> bool:
    """pandas' `is_scalar(value) and isna(value)` for the values a frame yields."""
    if value is None or type(value).__name__ in ("NaTType", "NAType"):
        return True
    if isinstance(value, decimal.Decimal):
        return value.is_nan()
    if isinstance(value, numbers.Real) and not isinstance(value, numbers.Integral):
        return value != value
    return False


def _list_like(value: Any) -> bool:
    return hasattr(value, "__iter__") and not isinstance(value, (str, bytes, type))


class _Handles:
    """The binary handle a writer saves to, and whether closing it is ours to do."""

    def __init__(self, path: Any, mode: str, storage_options: Any) -> None:
        self.created = False
        if isinstance(path, ExcelWriter) or hasattr(path, "write"):
            if storage_options is not None:
                raise ValueError(_STORAGE)
            self.handle = path
            return
        path = os.fspath(path)
        if isinstance(path, bytes):
            path = path.decode()
        if "://" in path:
            fsspec = _optional.imported("fsspec")
            self.handle = fsspec.open(path, mode=mode, **(storage_options or {})).open()
            self.created = True
            return
        if storage_options is not None:
            raise ValueError(_STORAGE)
        path = os.path.expanduser(path)
        if "r" not in mode:
            parent = Path(path).parent
            if not parent.is_dir():
                raise OSError(rf"Cannot save file into a non-existent directory: '{parent}'")
        self.handle = open(path, mode)  # noqa: SIM115
        self.created = True

    def close(self) -> None:
        if self.created:
            self.handle.close()
        self.created = False


def _combine_kwargs(engine_kwargs: dict[str, Any] | None, kwargs: dict[str, Any]) -> dict:
    result = {} if engine_kwargs is None else engine_kwargs.copy()
    result.update(kwargs)
    return result


def _valid_freeze_panes(freeze_panes: Any) -> bool:
    if freeze_panes is not None:
        if len(freeze_panes) == 2 and all(isinstance(item, int) for item in freeze_panes):
            return True
        raise ValueError(
            "freeze_panes must be of form (row, column) where row and column are integers"
        )
    return False


def _default_engine(ext: str) -> str:
    writers = {"xlsx": "openpyxl", "xlsm": "openpyxl", "xlsb": "pyxlsb", "ods": "odf"}
    try:
        import xlsxwriter  # noqa: F401
    except ImportError:
        pass
    else:
        writers["xlsx"] = "xlsxwriter"
    return writers[ext]


class ExcelWriter:
    """pandas' `ExcelWriter`, which becomes the writer for its engine when made."""

    _engine = ""
    _supported_extensions: tuple[str, ...] = ()

    def __new__(
        cls,
        path: Any,
        engine: str | None = None,
        date_format: str | None = None,
        datetime_format: str | None = None,
        mode: str = "w",
        storage_options: Any = None,
        if_sheet_exists: str | None = None,
        engine_kwargs: dict | None = None,
    ) -> Any:
        if cls is ExcelWriter:
            from ._config import get_option

            if engine is None or (isinstance(engine, str) and engine == "auto"):
                ext = os.path.splitext(path)[-1][1:] if isinstance(path, str) else "xlsx"
                try:
                    engine = get_option(f"io.excel.{ext}.writer")
                    if engine == "auto":
                        engine = _default_engine(ext)
                except KeyError as err:
                    raise ValueError(f"No engine for filetype: '{ext}'") from err
            try:
                cls = _WRITERS[engine]
            except (KeyError, TypeError) as err:
                raise ValueError(f"No Excel writer '{engine}'") from err
        return object.__new__(cls)

    @property
    def supported_extensions(self) -> tuple[str, ...]:
        """The extensions this engine writes."""
        return self._supported_extensions

    @property
    def engine(self) -> str:
        """The name of the engine."""
        return self._engine

    @property
    def sheets(self) -> dict[str, Any]:
        """The sheets of the book by name."""
        raise NotImplementedError

    @property
    def book(self) -> Any:
        """The engine's workbook."""
        raise NotImplementedError

    def _write_cells(
        self,
        cells: Any,
        sheet_name: str | None = None,
        startrow: int = 0,
        startcol: int = 0,
        freeze_panes: tuple[int, int] | None = None,
        autofilter_range: str | None = None,
    ) -> None:
        raise NotImplementedError

    def _save(self) -> None:
        raise NotImplementedError

    def __init__(
        self,
        path: Any,
        engine: str | None = None,
        date_format: str | None = None,
        datetime_format: str | None = None,
        mode: str = "w",
        storage_options: Any = None,
        if_sheet_exists: str | None = None,
        engine_kwargs: dict[str, Any] | None = None,
    ) -> None:
        if isinstance(path, str):
            self.check_extension(os.path.splitext(path)[-1])
        if "b" not in mode:
            mode += "b"
        mode = mode.replace("a", "r+")
        if if_sheet_exists not in (None, "error", "new", "replace", "overlay"):
            raise ValueError(_IF_SHEET_EXISTS.format(if_sheet_exists))
        if if_sheet_exists and "r+" not in mode:
            raise ValueError("if_sheet_exists is only valid in append mode (mode='a')")
        if if_sheet_exists is None:
            if_sheet_exists = "error"
        self._if_sheet_exists = if_sheet_exists
        self._handles = _Handles(path, mode, storage_options)
        self._cur_sheet = None
        self._date_format = "YYYY-MM-DD" if date_format is None else date_format
        self._datetime_format = (
            "YYYY-MM-DD HH:MM:SS" if datetime_format is None else datetime_format
        )
        self._mode = mode

    @property
    def date_format(self) -> str:
        """The number format for a date cell."""
        return self._date_format

    @property
    def datetime_format(self) -> str:
        """The number format for a moment cell."""
        return self._datetime_format

    @property
    def if_sheet_exists(self) -> str:
        """What appending does with a sheet that is already there."""
        return self._if_sheet_exists

    def __fspath__(self) -> str:
        return getattr(self._handles.handle, "name", "")

    def _get_sheet_name(self, sheet_name: str | None) -> str:
        if sheet_name is None:
            sheet_name = self._cur_sheet
        if sheet_name is None:
            raise ValueError("Must pass explicit sheet_name or set _cur_sheet property")
        return sheet_name

    def _value_with_fmt(self, val: Any) -> tuple[Any, str | None]:
        fmt = None
        if _is_integer(val):
            val = int(val)
        elif _is_float(val):
            val = float(val)
        elif _is_bool(val):
            val = bool(val)
        elif isinstance(val, decimal.Decimal):
            val = decimal.Decimal(val)
        elif isinstance(val, dt.datetime):
            fmt = self._datetime_format
        elif isinstance(val, dt.date):
            fmt = self._date_format
        elif isinstance(val, dt.timedelta):
            val = val.total_seconds() / 86400
            fmt = "0"
        else:
            val = str(val)
            if len(val) > 32767:
                warnings.warn(
                    f"Cell contents too long ({len(val)}), truncated to 32767 characters",
                    UserWarning,
                    stacklevel=4,
                )
        return val, fmt

    @classmethod
    def check_extension(cls, ext: str) -> bool:
        """Whether the engine writes this extension, and pandas' error when it does not."""
        if ext.startswith("."):
            ext = ext[1:]
        if not any(ext in extension for extension in cls._supported_extensions):
            raise ValueError(f"Invalid extension for engine '{cls.engine}': '{ext}'")
        return True

    def __enter__(self) -> Any:
        return self

    def __exit__(self, *args: Any) -> None:
        self.close()

    def close(self) -> None:
        """Save the book and close the file if the writer opened it."""
        self._save()
        self._handles.close()


class OpenpyxlWriter(ExcelWriter):
    """pandas' writer for xlsx and xlsm through openpyxl."""

    _engine = "openpyxl"
    _supported_extensions = (".xlsx", ".xlsm")

    def __init__(
        self,
        path: Any,
        engine: str | None = None,
        date_format: str | None = None,
        datetime_format: str | None = None,
        mode: str = "w",
        storage_options: Any = None,
        if_sheet_exists: str | None = None,
        engine_kwargs: dict[str, Any] | None = None,
        **kwargs: Any,
    ) -> None:
        from openpyxl.workbook import Workbook

        engine_kwargs = _combine_kwargs(engine_kwargs, kwargs)
        # pandas passes neither format on for openpyxl, so both stay the defaults.
        super().__init__(
            path,
            mode=mode,
            storage_options=storage_options,
            if_sheet_exists=if_sheet_exists,
            engine_kwargs=engine_kwargs,
        )
        if "r+" in self._mode:
            from openpyxl import load_workbook

            try:
                self._book = load_workbook(self._handles.handle, **engine_kwargs)
            except TypeError:
                self._handles.handle.close()
                raise
            self._handles.handle.seek(0)
        else:
            try:
                self._book = Workbook(**engine_kwargs)
            except TypeError:
                self._handles.handle.close()
                raise
            if self.book.worksheets:
                self.book.remove(self.book.worksheets[0])

    @property
    def book(self) -> Any:
        """The openpyxl workbook."""
        return self._book

    @property
    def sheets(self) -> dict[str, Any]:
        """The sheets of the book by name."""
        return {name: self.book[name] for name in self.book.sheetnames}

    def _save(self) -> None:
        self.book.save(self._handles.handle)
        if "r+" in self._mode:
            self._handles.handle.truncate()

    @classmethod
    def _convert_to_style_kwargs(cls, style_dict: dict[str, Any]) -> dict[str, Any]:
        """
        Convert a style_dict to a set of kwargs suitable for initializing
        or updating-on-copy an openpyxl v2 style object.

        Parameters
        ----------
        style_dict : dict
            A dict with zero or more of the following keys (or their synonyms).
                'font'
                'fill'
                'border' ('borders')
                'alignment'
                'number_format'
                'protection'

        Returns
        -------
        style_kwargs : dict
            A dict with the same, normalized keys as ``style_dict`` but each
            value has been replaced with a native openpyxl style object of the
            appropriate class.
        """
        _style_key_map = {"borders": "border"}

        style_kwargs: dict[str, Any] = {}
        for k, v in style_dict.items():
            k = _style_key_map.get(k, k)
            _conv_to_x = getattr(cls, f"_convert_to_{k}", lambda x: None)
            new_v = _conv_to_x(v)
            if new_v:
                style_kwargs[k] = new_v

        return style_kwargs

    @classmethod
    def _convert_to_color(cls, color_spec):
        """
        Convert ``color_spec`` to an openpyxl v2 Color object.

        Parameters
        ----------
        color_spec : str, dict
            A 32-bit ARGB hex string, or a dict with zero or more of the
            following keys.
                'rgb'
                'indexed'
                'auto'
                'theme'
                'tint'
                'index'
                'type'

        Returns
        -------
        color : openpyxl.styles.Color
        """
        from openpyxl.styles import Color

        if isinstance(color_spec, str):
            return Color(color_spec)
        else:
            return Color(**color_spec)

    @classmethod
    def _convert_to_font(cls, font_dict):
        """
        Convert ``font_dict`` to an openpyxl v2 Font object.

        Parameters
        ----------
        font_dict : dict
            A dict with zero or more of the following keys (or their synonyms).
                'name'
                'size' ('sz')
                'bold' ('b')
                'italic' ('i')
                'underline' ('u')
                'strikethrough' ('strike')
                'color'
                'vertAlign' ('vertalign')
                'charset'
                'scheme'
                'family'
                'outline'
                'shadow'
                'condense'

        Returns
        -------
        font : openpyxl.styles.Font
        """
        from openpyxl.styles import Font

        _font_key_map = {
            "sz": "size",
            "b": "bold",
            "i": "italic",
            "u": "underline",
            "strike": "strikethrough",
            "vertalign": "vertAlign",
        }

        font_kwargs = {}
        for k, v in font_dict.items():
            k = _font_key_map.get(k, k)
            if k == "color":
                v = cls._convert_to_color(v)
            font_kwargs[k] = v

        return Font(**font_kwargs)

    @classmethod
    def _convert_to_stop(cls, stop_seq):
        """
        Convert ``stop_seq`` to a list of openpyxl v2 Color objects,
        suitable for initializing the ``GradientFill`` ``stop`` parameter.

        Parameters
        ----------
        stop_seq : iterable
            An iterable that yields objects suitable for consumption by
            ``_convert_to_color``.

        Returns
        -------
        stop : list of openpyxl.styles.Color
        """
        return map(cls._convert_to_color, stop_seq)

    @classmethod
    def _convert_to_fill(cls, fill_dict: dict[str, Any]) -> Any:
        """
        Convert ``fill_dict`` to an openpyxl v2 Fill object.

        Parameters
        ----------
        fill_dict : dict
            A dict with one or more of the following keys (or their synonyms),
                'fill_type' ('patternType', 'patterntype')
                'start_color' ('fgColor', 'fgcolor')
                'end_color' ('bgColor', 'bgcolor')
            or one or more of the following keys (or their synonyms).
                'type' ('fill_type')
                'degree'
                'left'
                'right'
                'top'
                'bottom'
                'stop'

        Returns
        -------
        fill : openpyxl.styles.Fill
        """
        from openpyxl.styles import (
            GradientFill,
            PatternFill,
        )

        _pattern_fill_key_map = {
            "patternType": "fill_type",
            "patterntype": "fill_type",
            "fgColor": "start_color",
            "fgcolor": "start_color",
            "bgColor": "end_color",
            "bgcolor": "end_color",
        }

        _gradient_fill_key_map = {"fill_type": "type"}

        pfill_kwargs = {}
        gfill_kwargs = {}
        for k, v in fill_dict.items():
            pk = _pattern_fill_key_map.get(k)
            gk = _gradient_fill_key_map.get(k)
            if pk in ["start_color", "end_color"]:
                v = cls._convert_to_color(v)
            if gk == "stop":
                v = cls._convert_to_stop(v)
            if pk:
                pfill_kwargs[pk] = v
            elif gk:
                gfill_kwargs[gk] = v
            else:
                pfill_kwargs[k] = v
                gfill_kwargs[k] = v

        try:
            return PatternFill(**pfill_kwargs)
        except TypeError:
            return GradientFill(**gfill_kwargs)

    @classmethod
    def _convert_to_side(cls, side_spec):
        """
        Convert ``side_spec`` to an openpyxl v2 Side object.

        Parameters
        ----------
        side_spec : str, dict
            A string specifying the border style, or a dict with zero or more
            of the following keys (or their synonyms).
                'style' ('border_style')
                'color'

        Returns
        -------
        side : openpyxl.styles.Side
        """
        from openpyxl.styles import Side

        _side_key_map = {"border_style": "style"}

        if isinstance(side_spec, str):
            return Side(style=side_spec)

        side_kwargs = {}
        for k, v in side_spec.items():
            k = _side_key_map.get(k, k)
            if k == "color":
                v = cls._convert_to_color(v)
            side_kwargs[k] = v

        return Side(**side_kwargs)

    @classmethod
    def _convert_to_border(cls, border_dict):
        """
        Convert ``border_dict`` to an openpyxl v2 Border object.

        Parameters
        ----------
        border_dict : dict
            A dict with zero or more of the following keys (or their synonyms).
                'left'
                'right'
                'top'
                'bottom'
                'diagonal'
                'diagonal_direction'
                'vertical'
                'horizontal'
                'diagonalUp' ('diagonalup')
                'diagonalDown' ('diagonaldown')
                'outline'

        Returns
        -------
        border : openpyxl.styles.Border
        """
        from openpyxl.styles import Border

        _border_key_map = {"diagonalup": "diagonalUp", "diagonaldown": "diagonalDown"}

        border_kwargs = {}
        for k, v in border_dict.items():
            k = _border_key_map.get(k, k)
            if k == "color":
                v = cls._convert_to_color(v)
            if k in ["left", "right", "top", "bottom", "diagonal"]:
                v = cls._convert_to_side(v)
            border_kwargs[k] = v

        return Border(**border_kwargs)

    @classmethod
    def _convert_to_alignment(cls, alignment_dict):
        """
        Convert ``alignment_dict`` to an openpyxl v2 Alignment object.

        Parameters
        ----------
        alignment_dict : dict
            A dict with zero or more of the following keys (or their synonyms).
                'horizontal'
                'vertical'
                'text_rotation'
                'wrap_text'
                'shrink_to_fit'
                'indent'
        Returns
        -------
        alignment : openpyxl.styles.Alignment
        """
        from openpyxl.styles import Alignment

        return Alignment(**alignment_dict)

    @classmethod
    def _convert_to_number_format(cls, number_format_dict):
        """
        Convert ``number_format_dict`` to an openpyxl v2.1.0 number format
        initializer.

        Parameters
        ----------
        number_format_dict : dict
            A dict with zero or more of the following keys.
                'format_code' : str

        Returns
        -------
        number_format : str
        """
        return number_format_dict["format_code"]

    @classmethod
    def _convert_to_protection(cls, protection_dict):
        """
        Convert ``protection_dict`` to an openpyxl v2 Protection object.

        Parameters
        ----------
        protection_dict : dict
            A dict with zero or more of the following keys.
                'locked'
                'hidden'

        Returns
        -------
        """
        from openpyxl.styles import Protection

        return Protection(**protection_dict)

    def _write_cells(
        self,
        cells: Any,
        sheet_name: str | None = None,
        startrow: int = 0,
        startcol: int = 0,
        freeze_panes: tuple[int, int] | None = None,
        autofilter_range: str | None = None,
    ) -> None:
        sheet_name = self._get_sheet_name(sheet_name)
        if sheet_name in self.sheets and self._if_sheet_exists != "new":
            if "r+" in self._mode:
                if self._if_sheet_exists == "replace":
                    old = self.sheets[sheet_name]
                    target = self.book.index(old)
                    del self.book[sheet_name]
                    wks = self.book.create_sheet(sheet_name, target)
                elif self._if_sheet_exists == "error":
                    raise ValueError(
                        f"Sheet '{sheet_name}' already exists and "
                        f"if_sheet_exists is set to 'error'."
                    )
                elif self._if_sheet_exists == "overlay":
                    wks = self.sheets[sheet_name]
                else:
                    raise ValueError(_IF_SHEET_EXISTS.format(self._if_sheet_exists))
            else:
                wks = self.sheets[sheet_name]
        else:
            wks = self.book.create_sheet()
            wks.title = sheet_name
        style_cache: dict[str, dict[str, Any]] = {}
        if _valid_freeze_panes(freeze_panes):
            assert freeze_panes is not None
            wks.freeze_panes = wks.cell(row=freeze_panes[0] + 1, column=freeze_panes[1] + 1)
        for cell in cells:
            xcell = wks.cell(row=startrow + cell.row + 1, column=startcol + cell.col + 1)
            xcell.value, fmt = self._value_with_fmt(cell.val)
            if fmt:
                xcell.number_format = fmt
            style_kwargs: dict[str, Any] | None = {}
            if cell.style:
                key = str(cell.style)
                style_kwargs = style_cache.get(key)
                if style_kwargs is None:
                    style_kwargs = self._convert_to_style_kwargs(cell.style)
                    style_cache[key] = style_kwargs
            if style_kwargs:
                for k, v in style_kwargs.items():
                    setattr(xcell, k, v)
            if cell.mergestart is not None and cell.mergeend is not None:
                wks.merge_cells(
                    start_row=startrow + cell.row + 1,
                    start_column=startcol + cell.col + 1,
                    end_column=startcol + cell.mergeend + 1,
                    end_row=startrow + cell.mergestart + 1,
                )
                # Only the top left cell of a merged range is kept, so each of the
                # others is given the style too, as pandas does.
                if style_kwargs:
                    first_row = startrow + cell.row + 1
                    first_col = startcol + cell.col + 1
                    for row in range(first_row, startrow + cell.mergestart + 2):
                        for col in range(first_col, startcol + cell.mergeend + 2):
                            if row == first_row and col == first_col:
                                continue
                            other = wks.cell(column=col, row=row)
                            for k, v in style_kwargs.items():
                                setattr(other, k, v)
        if autofilter_range:
            wks.auto_filter.ref = autofilter_range


class _XlsxStyler:
    # Map from openpyxl-oriented styles to flatter xlsxwriter representation
    # Ordering necessary for both determinism and because some are keyed by
    # prefixes of others.
    STYLE_MAPPING: ClassVar[dict[str, list[tuple[tuple[str, ...], str]]]] = {
        "font": [
            (("name",), "font_name"),
            (("sz",), "font_size"),
            (("size",), "font_size"),
            (("color", "rgb"), "font_color"),
            (("color",), "font_color"),
            (("b",), "bold"),
            (("bold",), "bold"),
            (("i",), "italic"),
            (("italic",), "italic"),
            (("u",), "underline"),
            (("underline",), "underline"),
            (("strike",), "font_strikeout"),
            (("vertAlign",), "font_script"),
            (("vertalign",), "font_script"),
        ],
        "number_format": [(("format_code",), "num_format"), ((), "num_format")],
        "protection": [(("locked",), "locked"), (("hidden",), "hidden")],
        "alignment": [
            (("horizontal",), "align"),
            (("vertical",), "valign"),
            (("text_rotation",), "rotation"),
            (("wrap_text",), "text_wrap"),
            (("indent",), "indent"),
            (("shrink_to_fit",), "shrink"),
        ],
        "fill": [
            (("patternType",), "pattern"),
            (("patterntype",), "pattern"),
            (("fill_type",), "pattern"),
            (("start_color", "rgb"), "fg_color"),
            (("fgColor", "rgb"), "fg_color"),
            (("fgcolor", "rgb"), "fg_color"),
            (("start_color",), "fg_color"),
            (("fgColor",), "fg_color"),
            (("fgcolor",), "fg_color"),
            (("end_color", "rgb"), "bg_color"),
            (("bgColor", "rgb"), "bg_color"),
            (("bgcolor", "rgb"), "bg_color"),
            (("end_color",), "bg_color"),
            (("bgColor",), "bg_color"),
            (("bgcolor",), "bg_color"),
        ],
        "border": [
            (("color", "rgb"), "border_color"),
            (("color",), "border_color"),
            (("style",), "border"),
            (("top", "color", "rgb"), "top_color"),
            (("top", "color"), "top_color"),
            (("top", "style"), "top"),
            (("top",), "top"),
            (("right", "color", "rgb"), "right_color"),
            (("right", "color"), "right_color"),
            (("right", "style"), "right"),
            (("right",), "right"),
            (("bottom", "color", "rgb"), "bottom_color"),
            (("bottom", "color"), "bottom_color"),
            (("bottom", "style"), "bottom"),
            (("bottom",), "bottom"),
            (("left", "color", "rgb"), "left_color"),
            (("left", "color"), "left_color"),
            (("left", "style"), "left"),
            (("left",), "left"),
        ],
    }

    @classmethod
    def convert(cls, style_dict, num_format_str=None) -> dict[str, Any]:
        """
        converts a style_dict to an xlsxwriter format dict

        Parameters
        ----------
        style_dict : style dictionary to convert
        num_format_str : optional number format string
        """
        # Create an XlsxWriter format object.
        props = {}

        if num_format_str is not None:
            props["num_format"] = num_format_str

        if style_dict is None:
            return props

        if "borders" in style_dict:
            style_dict = style_dict.copy()
            style_dict["border"] = style_dict.pop("borders")

        for style_group_key, style_group in style_dict.items():
            for src, dst in cls.STYLE_MAPPING.get(style_group_key, []):
                # src is a sequence of keys into a nested dict
                # dst is a flat key
                if dst in props:
                    continue
                v = style_group
                for k in src:
                    try:
                        v = v[k]
                    except (KeyError, TypeError):
                        break
                else:
                    props[dst] = v

        if isinstance(props.get("pattern"), str):
            # TODO: support other fill patterns
            props["pattern"] = 0 if props["pattern"] == "none" else 1

        for k in ["border", "top", "right", "bottom", "left"]:
            if isinstance(props.get(k), str):
                try:
                    props[k] = [
                        "none",
                        "thin",
                        "medium",
                        "dashed",
                        "dotted",
                        "thick",
                        "double",
                        "hair",
                        "mediumDashed",
                        "dashDot",
                        "mediumDashDot",
                        "dashDotDot",
                        "mediumDashDotDot",
                        "slantDashDot",
                    ].index(props[k])
                except ValueError:
                    props[k] = 2

        if isinstance(props.get("font_script"), str):
            props["font_script"] = ["baseline", "superscript", "subscript"].index(
                props["font_script"]
            )

        if isinstance(props.get("underline"), str):
            props["underline"] = {
                "none": 0,
                "single": 1,
                "double": 2,
                "singleAccounting": 33,
                "doubleAccounting": 34,
            }[props["underline"]]

        # GH 30107 - xlsxwriter uses different name
        if props.get("valign") == "center":
            props["valign"] = "vcenter"

        return props


class XlsxWriter(ExcelWriter):
    """pandas' writer for xlsx through xlsxwriter."""

    _engine = "xlsxwriter"
    _supported_extensions = (".xlsx",)

    def __init__(
        self,
        path: Any,
        engine: str | None = None,
        date_format: str | None = None,
        datetime_format: str | None = None,
        mode: str = "w",
        storage_options: Any = None,
        if_sheet_exists: str | None = None,
        engine_kwargs: dict[str, Any] | None = None,
        **kwargs: Any,
    ) -> None:
        from xlsxwriter import Workbook

        engine_kwargs = _combine_kwargs(engine_kwargs, kwargs)
        if mode == "a":
            raise ValueError("Append mode is not supported with xlsxwriter!")
        super().__init__(
            path,
            engine=engine,
            date_format=date_format,
            datetime_format=datetime_format,
            mode=mode,
            storage_options=storage_options,
            if_sheet_exists=if_sheet_exists,
            engine_kwargs=engine_kwargs,
        )
        try:
            self._book = Workbook(self._handles.handle, **engine_kwargs)
        except TypeError:
            self._handles.handle.close()
            raise

    @property
    def book(self) -> Any:
        """The xlsxwriter workbook."""
        return self._book

    @property
    def sheets(self) -> dict[str, Any]:
        """The sheets of the book by name."""
        return self.book.sheetnames

    def _save(self) -> None:
        self.book.close()

    def _write_cells(
        self,
        cells: Any,
        sheet_name: str | None = None,
        startrow: int = 0,
        startcol: int = 0,
        freeze_panes: tuple[int, int] | None = None,
        autofilter_range: str | None = None,
    ) -> None:
        sheet_name = self._get_sheet_name(sheet_name)
        wks = self.book.get_worksheet_by_name(sheet_name)
        if wks is None:
            wks = self.book.add_worksheet(sheet_name)
        styles: dict[str, Any] = {"null": None}
        if _valid_freeze_panes(freeze_panes):
            assert freeze_panes is not None
            wks.freeze_panes(*freeze_panes)
        for cell in cells:
            val, fmt = self._value_with_fmt(cell.val)
            key = json.dumps(cell.style)
            if fmt:
                key += fmt
            if key in styles:
                style = styles[key]
            else:
                style = self.book.add_format(_XlsxStyler.convert(cell.style, fmt))
                styles[key] = style
            if cell.mergestart is not None and cell.mergeend is not None:
                wks.merge_range(
                    startrow + cell.row,
                    startcol + cell.col,
                    startrow + cell.mergestart,
                    startcol + cell.mergeend,
                    val,
                    style,
                )
            else:
                wks.write(startrow + cell.row, startcol + cell.col, val, style)
        if autofilter_range:
            wks.autofilter(autofilter_range)


class ODSWriter(ExcelWriter):
    """pandas' writer for OpenDocument spreadsheets through odfpy."""

    _engine = "odf"
    _supported_extensions = (".ods",)

    def __init__(
        self,
        path: Any,
        engine: str | None = None,
        date_format: str | None = None,
        datetime_format: str | None = None,
        mode: str = "w",
        storage_options: Any = None,
        if_sheet_exists: str | None = None,
        engine_kwargs: dict[str, Any] | None = None,
        **kwargs: Any,
    ) -> None:
        from odf.opendocument import OpenDocumentSpreadsheet

        if mode == "a":
            raise ValueError("Append mode is not supported with odf!")
        engine_kwargs = _combine_kwargs(engine_kwargs, kwargs)
        self._book = OpenDocumentSpreadsheet(**engine_kwargs)
        self._style_dict: dict[str, str] = {}
        super().__init__(
            path,
            mode=mode,
            storage_options=storage_options,
            if_sheet_exists=if_sheet_exists,
            engine_kwargs=engine_kwargs,
        )

    @property
    def book(self) -> Any:
        """The odfpy document."""
        return self._book

    @property
    def sheets(self) -> dict[str, Any]:
        """The sheets of the book by name."""
        from odf.table import Table

        return {sheet.getAttribute("name"): sheet for sheet in self.book.getElementsByType(Table)}

    def _save(self) -> None:
        for sheet in self.sheets.values():
            self.book.spreadsheet.addElement(sheet)
        self.book.save(self._handles.handle)

    def _write_cells(
        self,
        cells: Any,
        sheet_name: str | None = None,
        startrow: int = 0,
        startcol: int = 0,
        freeze_panes: tuple[int, int] | None = None,
        autofilter_range: str | None = None,
    ) -> None:
        if autofilter_range:
            raise ValueError("Autofilter is not supported with odf!")
        from odf.table import Table, TableCell, TableRow
        from odf.text import P

        sheet_name = self._get_sheet_name(sheet_name)
        if sheet_name in self.sheets:
            wks = self.sheets[sheet_name]
        else:
            wks = Table(name=sheet_name)
            self.book.spreadsheet.addElement(wks)
        if _valid_freeze_panes(freeze_panes):
            assert freeze_panes is not None
            self._create_freeze_panes(sheet_name, freeze_panes)
        for _ in range(startrow):
            wks.addElement(TableRow())
        rows: defaultdict[int, Any] = defaultdict(TableRow)
        counts: defaultdict[int, int] = defaultdict(int)
        for cell in sorted(cells, key=lambda cell: (cell.row, cell.col)):
            if not counts[cell.row]:
                for _ in range(startcol):
                    rows[cell.row].addElement(TableCell())
            for _ in range(cell.col - counts[cell.row]):
                rows[cell.row].addElement(TableCell())
                counts[cell.row] += 1
            shown, table_cell = self._make_table_cell(cell)
            rows[cell.row].addElement(table_cell)
            counts[cell.row] += 1
            table_cell.addElement(P(text=shown))
        if len(rows) > 0:
            for number in range(max(rows.keys()) + 1):
                wks.addElement(rows[number])

    def _make_table_cell(self, cell: _Cell) -> tuple[Any, Any]:
        from odf.table import TableCell

        attributes: dict[str, int | str] = {}
        style_name = self._process_style(cell.style)
        if style_name is not None:
            attributes["stylename"] = style_name
        if cell.mergestart is not None and cell.mergeend is not None:
            attributes["numberrowsspanned"] = max(1, cell.mergestart)
            attributes["numbercolumnsspanned"] = cell.mergeend
        val, _ = self._value_with_fmt(cell.val)
        if isinstance(val, bool):
            return str(val).upper(), TableCell(
                valuetype="boolean", booleanvalue=str(val).lower(), attributes=attributes
            )
        if isinstance(val, dt.datetime):
            return val.strftime("%c"), TableCell(
                valuetype="date", datevalue=val.isoformat(), attributes=attributes
            )
        if isinstance(val, dt.date):
            value = f"{val.year}-{val.month:02d}-{val.day:02d}"
            return val.strftime("%x"), TableCell(
                valuetype="date", datevalue=value, attributes=attributes
            )
        if isinstance(val, str):
            return val, TableCell(valuetype="string", stringvalue=val, attributes=attributes)
        return val, TableCell(valuetype="float", value=val, attributes=attributes)

    def _process_style(self, style: dict[str, Any] | None) -> str | None:
        """pandas' OpenDocument style for a cell's style, named once for each distinct style."""
        from odf.style import ParagraphProperties, Style, TableCellProperties, TextProperties

        if style is None:
            return None
        style_key = json.dumps(style)
        if style_key in self._style_dict:
            return self._style_dict[style_key]
        name = f"pd{len(self._style_dict) + 1}"
        self._style_dict[style_key] = name
        odf_style = Style(name=name, family="table-cell")
        if "font" in style:
            font = style["font"]
            if font.get("bold", False):
                odf_style.addElement(TextProperties(fontweight="bold"))
        if "borders" in style:
            borders = style["borders"]
            for side, thickness in borders.items():
                thickness_translation = {"thin": "0.75pt solid #000000"}
                odf_style.addElement(
                    TableCellProperties(
                        attributes={f"border{side}": thickness_translation[thickness]}
                    )
                )
        if "alignment" in style:
            alignment = style["alignment"]
            horizontal = alignment.get("horizontal")
            if horizontal:
                odf_style.addElement(ParagraphProperties(textalign=horizontal))
            vertical = alignment.get("vertical")
            if vertical:
                odf_style.addElement(TableCellProperties(verticalalign=vertical))
        self.book.styles.addElement(odf_style)
        return name

    def _create_freeze_panes(self, sheet_name: str, freeze_panes: tuple[int, int]) -> None:
        from odf.config import (
            ConfigItem,
            ConfigItemMapEntry,
            ConfigItemMapIndexed,
            ConfigItemMapNamed,
            ConfigItemSet,
        )

        item_set = ConfigItemSet(name="ooo:view-settings")
        self.book.settings.addElement(item_set)
        indexed = ConfigItemMapIndexed(name="Views")
        item_set.addElement(indexed)
        entry = ConfigItemMapEntry()
        indexed.addElement(entry)
        named = ConfigItemMapNamed(name="Tables")
        entry.addElement(named)
        entry = ConfigItemMapEntry(name=sheet_name)
        named.addElement(entry)
        for name, kind, text in (
            ("HorizontalSplitMode", "short", "2"),
            ("VerticalSplitMode", "short", "2"),
            ("HorizontalSplitPosition", "int", str(freeze_panes[0])),
            ("VerticalSplitPosition", "int", str(freeze_panes[1])),
            ("PositionRight", "int", str(freeze_panes[0])),
            ("PositionBottom", "int", str(freeze_panes[1])),
        ):
            entry.addElement(ConfigItem(name=name, type=kind, text=text))


for _writer in (OpenpyxlWriter, XlsxWriter, ODSWriter):
    _WRITERS[_writer._engine] = _writer


class _Cell:
    """One cell for a writer: where it goes, what it holds, and the span it merges."""

    def __init__(
        self,
        row: int,
        col: int,
        val: Any,
        mergestart: int | None = None,
        mergeend: int | None = None,
        style: Any = None,
    ) -> None:
        self.row = row
        self.col = col
        self.val = val
        self.style = style
        self.mergestart = mergestart
        self.mergeend = mergeend


def _formatted(value: Any) -> str:
    return "NaN" if _missing(value) else str(value)


def _sparsified(levels: list[list[str]]) -> list[list[str]]:
    """pandas' `sparsify_labels`: a label repeated with every level above it is blanked."""
    rows = list(zip(*levels, strict=True))
    if not rows:
        return levels
    count = len(levels)
    result = [list(rows[0])]
    for prev, cur in itertools.pairwise(rows):
        sparse: list[str] = []
        for i, (p, t) in enumerate(zip(prev, cur, strict=True)):
            if i == count - 1:
                sparse.append(t)
                break
            if p == t:
                sparse.append("")
            else:
                sparse.extend(cur[i:])
                break
        result.append(sparse)
    return [list(level) for level in zip(*result, strict=True)]


def _level_lengths(levels: list[list[str]]) -> list[dict[int, int]]:
    """pandas' `get_level_lengths`: where each label starts, and how many rows it spans."""
    if len(levels) == 0:
        return []
    control = [True] * len(levels[0])
    result = []
    for level in levels:
        last = 0
        lengths = {}
        for i, key in enumerate(level):
            if control[i] and key == "":
                pass
            else:
                control[i] = False
                lengths[last] = i - last
                last = i
        lengths[last] = len(level) - last
        result.append(lengths)
    return result


def _multi_levels(index: Any, sparsify: bool) -> tuple[list[dict[int, int]], list[list[Any]]]:
    values = [list(index.get_level_values(i)) for i in range(index.nlevels)]
    shown = [[_formatted(v) for v in level] for level in values]
    if sparsify:
        shown = _sparsified(shown)
    return _level_lengths(shown), values


def _timestamp(value: Any) -> Any:
    from ._period import Period

    return value.to_timestamp() if isinstance(value, Period) else value


class _Formatter:
    """pandas' `ExcelFormatter` for a frame: the cells of the header, the index and the body."""

    max_rows = 2**20
    max_cols = 2**14

    def __init__(
        self,
        df: Any,
        na_rep: str = "",
        float_format: str | None = None,
        cols: Any = None,
        header: Any = True,
        index: bool = True,
        index_label: Any = None,
        merge_cells: Any = False,
        inf_rep: str = "inf",
        style_converter: Any = None,
        autofilter: bool = False,
    ) -> None:
        from ._frame import DataFrame

        self.rowcounter = 0
        self.na_rep = na_rep
        self.styler: Any = None
        self.style_converter: Any = None
        if not isinstance(df, DataFrame):
            from ._css import CSSToExcelConverter

            self.styler = df
            self.styler._compute()
            df = df.data
            self.style_converter = style_converter or CSSToExcelConverter()
        self.df = df
        if cols is not None:
            if not _list_like(cols):
                raise TypeError(
                    f"Index(...) must be called with a collection of some kind, {cols!r} was passed"
                )
            present = set(cols) & set(df.columns)
            if not present:
                raise KeyError("passes columns are not ALL present dataframe")
            if len(present) != len(set(cols)):
                raise KeyError("Not all names specified in 'columns' are found")
            self.df = df.reindex(columns=list(cols))
        self.columns = self.df.columns
        self.float_format = float_format
        self.index = index
        self.index_label = index_label
        self.header = header
        if not isinstance(merge_cells, bool) and merge_cells != "columns":
            raise ValueError(f"Unexpected value for merge_cells={merge_cells!r}.")
        self.merge_cells = merge_cells
        self.inf_rep = inf_rep
        self.autofilter = autofilter

    def _styled(self, table: str, at: tuple[int, int], *cell: Any) -> _Cell:
        """pandas' `CssExcelCell`: a cell styled from the Styler's CSS for one position."""
        styles = getattr(self.styler, table, None)
        style = None
        if styles and self.style_converter:
            declarations = {prop.lower(): val for prop, val in styles[at]}
            style = self.style_converter(frozenset(declarations.items()))
        return _Cell(*cell, style=style)

    @staticmethod
    def _multi(index: Any) -> bool:
        from ._multi import MultiIndex

        return isinstance(index, MultiIndex)

    def _format_value(self, val: Any) -> Any:
        if _missing(val):
            val = self.na_rep
        elif _is_float(val):
            if val == math.inf:
                val = self.inf_rep
            elif val == -math.inf:
                val = f"-{self.inf_rep}"
            elif self.float_format is not None:
                val = float(self.float_format % val)
        if getattr(val, "tzinfo", None) is not None:
            raise ValueError(
                "Excel does not support datetimes with "
                "timezones. Please ensure that datetimes "
                "are timezone unaware before writing to Excel."
            )
        return val

    def _format_header_mi(self) -> Iterator[_Cell]:
        if self.columns.nlevels > 1 and not self.index:
            raise NotImplementedError(
                "Writing to Excel with MultiIndex columns and no "
                "index ('index'=False) is not yet implemented."
            )
        if not (self._has_aliases or self.header):
            return
        columns = self.columns
        merge_columns = self.merge_cells in {True, "columns"}
        level_lengths, level_values = _multi_levels(columns, merge_columns)
        coloffset = 0
        lnum = 0
        if self.index and self._multi(self.df.index):
            coloffset = self.df.index.nlevels - 1
        for lnum, name in enumerate(columns.names):
            yield _Cell(lnum, coloffset, name)
        for lnum, (spans, values) in enumerate(zip(level_lengths, level_values, strict=True)):
            for i, span in spans.items():
                start, end = None, None
                if merge_columns and span > 1:
                    start, end = lnum, coloffset + i + span
                yield self._styled(
                    "ctx_columns", (lnum, i), lnum, coloffset + i + 1, values[i], start, end
                )
        self.rowcounter = lnum

    def _format_header_regular(self) -> Iterator[_Cell]:
        if self._has_aliases or self.header:
            coloffset = 0
            if self.index:
                coloffset = 1
                if self._multi(self.df.index):
                    coloffset = len(self.df.index.names)
            names = list(self.columns)
            if self._has_aliases:
                if len(self.header) != len(self.columns):
                    raise ValueError(
                        f"Writing {len(self.columns)} cols but got {len(self.header)} aliases"
                    )
                names = self.header
            for at, name in enumerate(names):
                yield self._styled("ctx_columns", (0, at), self.rowcounter, at + coloffset, name)

    def _format_header(self) -> Iterator[_Cell]:
        if self._multi(self.columns):
            gen: Any = self._format_header_mi()
        else:
            gen = self._format_header_regular()
        names: Any = ()
        if self.df.index.names:
            row = [x if x is not None else "" for x in self.df.index.names]
            row += [""] * len(self.columns)
            if functools.reduce(lambda x, y: x and y, (x != "" for x in row)):
                names = (_Cell(self.rowcounter, at, val) for at, val in enumerate(row))
                self.rowcounter += 1
        return itertools.chain(gen, names)

    def _format_body(self) -> Iterator[_Cell]:
        if self._multi(self.df.index):
            return self._format_hierarchical_rows()
        return self._format_regular_rows()

    def _format_regular_rows(self) -> Iterator[_Cell]:
        from ._period_index import PeriodIndex

        if self._has_aliases or self.header:
            self.rowcounter += 1
        if self.index:
            if self.index_label and _list_like(self.index_label):
                index_label = self.index_label[0]
            elif self.index_label and isinstance(self.index_label, str):
                index_label = self.index_label
            else:
                index_label = self.df.index.names[0]
            if self._multi(self.columns):
                self.rowcounter += 1
            if index_label and self.header is not False:
                yield _Cell(self.rowcounter - 1, 0, index_label)
            values = self.df.index
            if isinstance(values, PeriodIndex):
                values = values.to_timestamp()
            for at, value in enumerate(values):
                yield self._styled("ctx_index", (at, 0), self.rowcounter + at, 0, value)
            coloffset = 1
        else:
            coloffset = 0
        yield from self._generate_body(coloffset)

    def _format_hierarchical_rows(self) -> Iterator[_Cell]:
        if self._has_aliases or self.header:
            self.rowcounter += 1
        gcolidx = 0
        if self.index:
            index_labels = list(self.df.index.names)
            if self.index_label and _list_like(self.index_label):
                index_labels = list(self.index_label)
            if self._multi(self.columns):
                self.rowcounter += 1
            if any(name is not None for name in index_labels) and self.header is not False:
                for at, name in enumerate(index_labels):
                    yield _Cell(self.rowcounter - 1, at, name)
            if self.merge_cells and self.merge_cells != "columns":
                level_lengths, level_values = _multi_levels(self.df.index, True)
                for spans, values in zip(level_lengths, level_values, strict=False):
                    for i, span in spans.items():
                        start, end = None, None
                        if span > 1:
                            start = self.rowcounter + i + span - 1
                            end = gcolidx
                        yield self._styled(
                            "ctx_index",
                            (i, gcolidx),
                            self.rowcounter + i,
                            gcolidx,
                            _timestamp(values[i]),
                            start,
                            end,
                        )
                    gcolidx += 1
            else:
                for level in zip(*self.df.index, strict=True):
                    for at, value in enumerate(level):
                        yield self._styled(
                            "ctx_index",
                            (at, gcolidx),
                            self.rowcounter + at,
                            gcolidx,
                            _timestamp(value),
                        )
                    gcolidx += 1
        yield from self._generate_body(gcolidx)

    @property
    def _has_aliases(self) -> bool:
        return _list_like(self.header)

    def _generate_body(self, coloffset: int) -> Iterator[_Cell]:
        for at in range(len(self.columns)):
            for i, val in enumerate(self.df.iloc[:, at]):
                yield self._styled("ctx", (i, at), self.rowcounter + i, at + coloffset, val)

    def get_formatted_cells(self) -> Iterator[_Cell]:
        for cell in itertools.chain(self._format_header(), self._format_body()):
            cell.val = self._format_value(cell.val)
            yield cell

    @staticmethod
    def _num2excel(index: int) -> str:
        if index < 0:
            raise ValueError(f"Index cannot be negative: {index}")
        name = ""
        while index > 0 or not name:
            index, remainder = divmod(index, 26)
            name = chr(65 + remainder) + name
        return name

    def write(
        self,
        writer: Any,
        sheet_name: str = "Sheet1",
        startrow: int = 0,
        startcol: int = 0,
        freeze_panes: tuple[int, int] | None = None,
        engine: str | None = None,
        storage_options: Any = None,
        engine_kwargs: dict | None = None,
    ) -> None:
        num_rows, num_cols = self.df.shape
        if num_rows > self.max_rows or num_cols > self.max_cols:
            raise ValueError(
                f"This sheet is too large! Your sheet size is: {num_rows}, {num_cols} "
                f"Max sheet size is: {self.max_rows}, {self.max_cols}"
            )
        autofilter_range = None
        if self.autofilter:
            startrowsoffset = 1
            endrowsoffset = 1
            merged = (
                "Excel filters merged cells by showing only the first row. "
                "'autofilter' and 'merge_cells' cannot be used simultaneously."
            )
            if num_cols == 0:
                indexoffset = 0
            elif self.index:
                indexoffset = 0
                if self._multi(self.df.index):
                    if self.merge_cells:
                        raise ValueError(merged)
                    indexoffset = self.df.index.nlevels - 1
                if self._multi(self.columns):
                    if self.merge_cells:
                        raise ValueError(merged)
                    startrowsoffset = self.columns.nlevels
                    endrowsoffset = self.columns.nlevels + 1
            else:
                indexoffset = -1
            start = f"{self._num2excel(startcol)}{startrow + startrowsoffset}"
            end_column = self._num2excel(startcol + num_cols + indexoffset)
            end = f"{end_column}{startrow + num_rows + endrowsoffset}"
            autofilter_range = f"{start}:{end}"
        if engine_kwargs is None:
            engine_kwargs = {}
        cells = self.get_formatted_cells()
        if isinstance(writer, ExcelWriter):
            need_save = False
        else:
            writer = ExcelWriter(
                writer,
                engine=engine,
                storage_options=storage_options,
                engine_kwargs=engine_kwargs,
            )
            need_save = True
        try:
            writer._write_cells(
                cells,
                sheet_name,
                startrow=startrow,
                startcol=startcol,
                freeze_panes=freeze_panes,
                autofilter_range=autofilter_range,
            )
        finally:
            if need_save:
                writer.close()


def to_excel(
    self: Any,
    excel_writer: Any,
    *,
    sheet_name: str = "Sheet1",
    na_rep: str = "",
    float_format: str | None = None,
    columns: Any = None,
    header: Any = True,
    index: bool = True,
    index_label: Any = None,
    startrow: int = 0,
    startcol: int = 0,
    engine: str | None = None,
    merge_cells: bool = True,
    inf_rep: str = "inf",
    freeze_panes: tuple[int, int] | None = None,
    storage_options: Any = None,
    engine_kwargs: dict[str, Any] | None = None,
    autofilter: bool = False,
) -> None:
    """Write the frame, or a series as a one column frame, to a sheet of a workbook."""
    if engine_kwargs is None:
        engine_kwargs = {}
    df = self if self.ndim == 2 else self.to_frame()
    formatter = _Formatter(
        df,
        na_rep=na_rep,
        cols=columns,
        header=header,
        float_format=float_format,
        index=index,
        index_label=index_label,
        merge_cells=merge_cells,
        inf_rep=inf_rep,
        autofilter=autofilter,
    )
    formatter.write(
        excel_writer,
        sheet_name=sheet_name,
        startrow=startrow,
        startcol=startcol,
        freeze_panes=freeze_panes,
        engine=engine,
        storage_options=storage_options,
        engine_kwargs=engine_kwargs,
    )
