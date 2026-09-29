"""`to_excel` and `ExcelWriter`, compared with pandas.

Each test writes the same frame with both libraries and the same engine, then
opens the two books and compares what they hold: every cell's value and number
format, the merged ranges, the frozen panes and the filter for xlsx, and the
sheet's XML for OpenDocument. Nothing in a saved book depends on which
library wrote it, so any difference is a difference in the cells.
"""

from __future__ import annotations

import datetime as dt
import importlib.util
import inspect
import io
import re
import zipfile
from collections.abc import Callable
from pathlib import Path
from types import ModuleType
from typing import Any

import pytest

import firepanda as fp

needs_pandas = pytest.mark.skipif(
    importlib.util.find_spec("pandas") is None, reason="pandas is not installed"
)
needs_openpyxl = pytest.mark.skipif(
    importlib.util.find_spec("openpyxl") is None, reason="openpyxl is not installed"
)


@needs_pandas
@pytest.mark.parametrize("path", ["DataFrame.to_excel", "Series.to_excel", "ExcelWriter"])
def test_the_signature_is_pandas(path: str) -> None:
    import pandas

    def found(lib: ModuleType) -> Any:
        target: Any = lib
        for part in path.split("."):
            target = getattr(target, part)
        return target

    ours = inspect.signature(found(fp))
    theirs = inspect.signature(found(pandas))
    assert list(ours.parameters) == list(theirs.parameters)


def book_of(data: bytes) -> Any:
    """What an xlsx book holds, read back with openpyxl."""
    import openpyxl

    book = openpyxl.load_workbook(io.BytesIO(data))
    return [
        (
            sheet.title,
            [[(cell.value, cell.number_format) for cell in row] for row in sheet.iter_rows()],
            sorted(str(merged) for merged in sheet.merged_cells.ranges),
            sheet.freeze_panes,
            sheet.auto_filter.ref,
        )
        for sheet in book.worksheets
    ]


def content_of(data: bytes) -> str:
    """The sheets of an OpenDocument book, without the namespaces odfpy happened to load."""
    xml = zipfile.ZipFile(io.BytesIO(data)).read("content.xml").decode()
    return re.sub(r' (xmlns:\w+|office:version)="[^"]*"', "", xml)


def written(lib: ModuleType, build: Callable[[ModuleType], Any], engine: str, options: Any) -> Any:
    buf = io.BytesIO()
    try:
        build(lib).to_excel(buf, engine=engine, **options)
    except Exception as err:
        return "raise", type(err).__name__, str(err)
    return content_of(buf.getvalue()) if engine == "odf" else book_of(buf.getvalue())


BUILT: list[Callable[[ModuleType], Any]] = [
    lambda m: m.DataFrame(
        {
            "i": [1, 2, 3],
            "f": [1.5, None, float("inf")],
            "s": ["a", None, "c"],
            "b": [True, False, True],
        }
    ),
    lambda m: m.DataFrame(
        {
            "t": m.to_datetime(["2020-01-01 10:00", None, "2021-03-04 00:00"]),
            "d": m.to_timedelta([1, None, 3600], unit="s"),
        }
    ),
    lambda m: m.DataFrame({"k": ["x", "y"], "v": [1, 2]}).set_index("k"),
    lambda m: m.DataFrame(
        {"k": ["x", "x", "y", "y"], "j": [1, 2, 1, 1], "v": [1.0, 2.0, 3.0, 4.0], "w": [5, 6, 7, 8]}
    ).set_index(["k", "j"]),
    # The column levels are left unnamed, because firepanda does not keep the
    # names of a column MultiIndex yet.
    lambda m: m.DataFrame([[1, 2, 3], [4, 5, 6]], index=["r", "s"]).set_axis(
        m.MultiIndex.from_tuples([("x", 1), ("x", 2), ("y", 1)]), axis=1
    ),
    lambda m: m.Series([1.5, 2.5], index=["p", "q"], name="n"),
    lambda m: m.Series([1, 2]),
    lambda m: m.DataFrame({"v": [1, 2]}, index=m.period_range("2020-01", periods=2, freq="M")),
    lambda m: m.DataFrame({"z": m.Series(m.to_datetime(["2020-01-01"])).dt.tz_localize("UTC")}),
]

OPTIONS: list[dict[str, Any]] = [
    {},
    {"index": False},
    {"header": False},
    {"merge_cells": False},
    {"merge_cells": "columns"},
    {"na_rep": "NA", "inf_rep": "INF"},
    {"float_format": "%.1f"},
    {"startrow": 2, "startcol": 1},
    {"freeze_panes": (1, 1)},
    {"autofilter": True},
    {"index_label": "L"},
    {"index_label": ["A", "B"]},
    {"header": ["h1", "h2", "h3", "h4"]},
    {"columns": ["i", "s"]},
    {"columns": ["zz"]},
    {"sheet_name": "Other"},
    {"merge_cells": "bad"},
    {"freeze_panes": (1,)},
]


def engines() -> list[Any]:
    marks = {"openpyxl": "openpyxl", "xlsxwriter": "xlsxwriter", "odf": "odf"}
    return [
        pytest.param(
            engine,
            marks=pytest.mark.skipif(
                importlib.util.find_spec(module) is None
                or importlib.util.find_spec("openpyxl") is None,
                reason=f"{module} is not installed",
            ),
        )
        for engine, module in marks.items()
    ]


@needs_pandas
@pytest.mark.parametrize("engine", engines())
@pytest.mark.parametrize("build", BUILT)
@pytest.mark.parametrize("options", OPTIONS)
def test_the_book_is_pandas_book(
    engine: str, build: Callable[[ModuleType], Any], options: dict[str, Any]
) -> None:
    import pandas

    assert written(fp, build, engine, options) == written(pandas, build, engine, options)


def frame(m: ModuleType) -> Any:
    return m.DataFrame({"a": [1, 2], "t": [dt.datetime(2020, 1, 2, 3, 4), dt.date(2020, 5, 6)]})


@needs_pandas
@pytest.mark.parametrize("engine", engines()[:2])
def test_a_writer_holds_several_sheets(engine: str) -> None:
    import pandas

    def made(m: ModuleType) -> Any:
        buf = io.BytesIO()
        with m.ExcelWriter(
            buf, engine=engine, date_format="DD/MM/YYYY", datetime_format="YYYY HH:MM"
        ) as writer:
            frame(m).to_excel(writer, sheet_name="A")
            frame(m).to_excel(writer, sheet_name="B", startrow=3, index=False)
            frame(m).to_excel(writer, sheet_name="A", startcol=5)
            about = (
                writer.engine,
                type(writer).__name__,
                sorted(writer.sheets),
                writer.date_format,
                writer.datetime_format,
                writer.if_sheet_exists,
                writer.supported_extensions,
            )
        return about, book_of(buf.getvalue())

    assert made(fp) == made(pandas)


@needs_pandas
@needs_openpyxl
@pytest.mark.parametrize("if_sheet_exists", [None, "error", "new", "replace", "overlay"])
def test_appending_is_pandas_appending(tmp_path: Path, if_sheet_exists: Any) -> None:
    import pandas

    def made(m: ModuleType) -> Any:
        path = tmp_path / f"{m.__name__}.xlsx"
        frame(m).to_excel(path, engine="openpyxl")
        try:
            with m.ExcelWriter(
                path, mode="a", engine="openpyxl", if_sheet_exists=if_sheet_exists
            ) as writer:
                frame(m).to_excel(writer, sheet_name="Sheet1", startrow=5)
                frame(m).to_excel(writer, sheet_name="New")
        except ValueError as err:
            return str(err)
        return book_of(path.read_bytes())

    assert made(fp) == made(pandas)


MISTAKES: list[Callable[[ModuleType, Path], Any]] = [
    lambda m, p: frame(m).to_excel(str(p / "a.csv")),
    lambda m, p: frame(m).to_excel(str(p / "a")),
    lambda m, p: frame(m).to_excel(io.BytesIO(), engine="nope"),
    lambda m, p: m.ExcelWriter(str(p / "a.xlsx"), engine="xlsxwriter", mode="a"),
    lambda m, p: m.ExcelWriter(str(p / "a.ods"), engine="odf", mode="a"),
    lambda m, p: m.ExcelWriter(io.BytesIO(), engine="openpyxl", if_sheet_exists="x"),
    lambda m, p: m.ExcelWriter(io.BytesIO(), engine="openpyxl", if_sheet_exists="new"),
    lambda m, p: frame(m).to_excel(str(p / "nodir" / "a.xlsx")),
    lambda m, p: frame(m).to_excel(io.BytesIO(), storage_options={"a": 1}),
    lambda m, p: frame(m).to_excel(io.BytesIO(), engine="odf", autofilter=True),
    lambda m, p: frame(m).to_excel(io.BytesIO(), header=["x"]),
    lambda m, p: frame(m).to_excel(io.BytesIO(), columns="a"),
    lambda m, p: frame(m).to_excel(io.BytesIO(), columns=["a", "zz"]),
    lambda m, p: frame(m).to_excel(io.BytesIO(), engine="openpyxl", engine_kwargs={"bogus": 1}),
    lambda m, p: (
        m.DataFrame([[1, 2]])
        .set_axis(m.MultiIndex.from_tuples([("a", 1), ("b", 2)]), axis=1)
        .to_excel(io.BytesIO(), index=False)
    ),
]


@needs_pandas
@needs_openpyxl
@pytest.mark.parametrize("call", MISTAKES)
def test_a_mistake_is_pandas_mistake(
    tmp_path: Path, call: Callable[[ModuleType, Path], Any]
) -> None:
    import pandas

    with pytest.raises(Exception) as theirs:
        call(pandas, tmp_path)
    with pytest.raises(Exception) as ours:
        call(fp, tmp_path)
    assert type(ours.value) is type(theirs.value)
    assert str(ours.value) == str(theirs.value)


@needs_openpyxl
def test_a_book_reads_back(tmp_path: Path) -> None:
    path = tmp_path / "book.xlsx"
    written = fp.DataFrame({"a": [1, 2], "b": ["x", None], "c": [1.5, 2.5]}, index=["p", "q"])
    written.to_excel(path)
    back = fp.read_excel(path, index_col=0)
    assert back.to_dict() == written.to_dict()
    assert list(back.index) == ["p", "q"]
