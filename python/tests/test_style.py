"""`DataFrame.style`, checked against pandas.

A frame's `style` is a `Styler` that writes the frame as HTML, LaTeX, text or
Typst, with the formats, colors, bars, tooltips and hidden rows pandas' Styler
takes, and each table here is compared with pandas' own text for the same
calls. Without jinja2 the accessor is refused with pandas' AttributeError.
"""

from __future__ import annotations

import importlib.util
import re
import zipfile
from collections.abc import Callable
from types import ModuleType
from typing import Any

import pytest

both = pytest.mark.skipif(
    importlib.util.find_spec("pandas") is None or importlib.util.find_spec("jinja2") is None,
    reason="pandas and jinja2 are not both installed",
)

H = {"table_uuid": "u"}


def mixed(m: ModuleType) -> Any:
    """Floats with a gap, whole numbers and text, with text row labels."""
    return m.DataFrame(
        {"a": [1.5, -2.0, None, 4.25], "b": [1, 2, 3, 4], "s": ["x", "y", None, "z"]},
        index=["r0", "r1", "r2", "r3"],
    )


def levels(m: ModuleType) -> Any:
    """Rows and columns labelled by two levels."""
    index = m.MultiIndex.from_tuples([("A", 1), ("A", 2), ("B", 1), ("B", 2)], names=["k", "n"])
    columns = m.MultiIndex.from_tuples([("x", "p"), ("x", "q"), ("y", "p")])
    rows = [[1, 2.5, 3], [4, 5.5, 6], [7, 8.5, 9], [10, 11.5, 12]]
    return m.DataFrame(rows, index=index, columns=columns)


def small(m: ModuleType) -> Any:
    """Two columns of two whole numbers."""
    return m.DataFrame({"a": [1, -5], "b": [2, 6]})


def tips(m: ModuleType) -> Any:
    """Tooltip text for the first column of `mixed`."""
    return m.DataFrame({"a": ["t0", None, "t2", ""]}, index=["r0", "r1", "r2", "r3"])


def bold_first(column: Any) -> list[str]:
    """Bold for the first cell of a column."""
    return ["font-weight: bold" if at == 0 else "" for at in range(len(column))]


TABLES: list[Callable[[Any], Any]] = [
    lambda m: mixed(m).style.to_html(**H),
    lambda m: levels(m).style.to_html(**H),
    lambda m: levels(m).style.to_html(**H, sparse_index=False, sparse_columns=False),
    lambda m: mixed(m).style.format(precision=2, na_rep="-", thousands=",").to_html(**H),
    lambda m: mixed(m).style.format({"a": "{:.1f}", "s": str.upper}, na_rep="NA").to_html(**H),
    lambda m: mixed(m).style.highlight_max(subset=["a", "b"]).to_html(**H),
    lambda m: mixed(m).style.highlight_min(axis=1, subset=["a", "b"]).to_html(**H),
    lambda m: mixed(m).style.highlight_max(axis=None, subset=["a", "b"]).to_html(**H),
    lambda m: mixed(m).style.highlight_null().to_html(**H),
    lambda m: mixed(m).style.highlight_between(left=0, right=3, subset=["a", "b"]).to_html(**H),
    lambda m: mixed(m).style.highlight_quantile(q_left=0.5, subset=["b"]).to_html(**H),
    lambda m: small(m).style.highlight_between(left=[0, 3], right=[2, 9], axis=1).to_html(**H),
    lambda m: mixed(m).style.bar().to_html(**H),
    lambda m: mixed(m).style.bar(align="zero", color=["red", "green"], height=50).to_html(**H),
    lambda m: small(m).style.bar(align="mean", axis=None, vmax=4).to_html(**H),
    lambda m: small(m).style.bar(align=2, axis=1).to_html(**H),
    lambda m: m.DataFrame({"a": [-1, -5]}).style.bar().to_html(**H),
    lambda m: (
        mixed(m).style.map(lambda v: "color: red;" if isinstance(v, str) else "").to_html(**H)
    ),
    lambda m: mixed(m).style.apply(bold_first).to_html(**H),
    lambda m: (
        mixed(m)
        .style.apply(lambda r: ["color: blue"] * len(r), axis=1, subset=(["r0", "r1"], ["a", "b"]))
        .to_html(**H)
    ),
    lambda m: mixed(m).style.hide(["r1"]).hide(["b"], axis=1).to_html(**H),
    lambda m: levels(m).style.hide(level="n").hide(axis=1, level=0).to_html(**H),
    lambda m: (
        mixed(m)
        .style.set_properties(subset=["a"], color="red")
        .set_caption("cap")
        .set_table_styles([{"selector": "th", "props": "color: blue;"}])
        .to_html(**H)
    ),
    lambda m: levels(m).style.set_sticky(axis=1).set_sticky(axis=0).to_html(**H),
    lambda m: mixed(m).style.set_tooltips(tips(m)).to_html(**H),
    lambda m: mixed(m).style.set_tooltips(tips(m), as_title_attribute=True).to_html(**H),
    lambda m: mixed(m).style.to_html(**H, max_rows=2, max_columns=2),
    lambda m: (
        levels(m)
        .style.format_index(str.lower, axis=1)
        .relabel_index(["a", "b", "c", "d"], level=1)
        .to_html(**H)
    ),
    lambda m: (
        levels(m)
        .style.map_index(lambda v: "color: red;" if v == "A" else "", level=0)
        .apply_index(lambda s: ["color: blue;"] * len(s), axis=1)
        .to_html(**H)
    ),
    lambda m: (
        mixed(m)[["a", "b"]].style.concat(mixed(m)[["a", "b"]].agg(["sum"]).style).to_html(**H)
    ),
    lambda m: mixed(m).style.use(mixed(m).style.highlight_max(subset=["b"]).export()).to_html(**H),
    lambda m: m.DataFrame({"<a>": ["<b>&", "c_%"]}).style.format(escape="html").to_html(**H),
    lambda m: (
        m.DataFrame({"a": ["see https://x.org/p now"]}).style.format(hyperlinks="html").to_html(**H)
    ),
    lambda m: (
        m.DataFrame({"a": [1]}, index=m.Index([5], name="k")).style.hide(names=True).to_html(**H)
    ),
    lambda m: m.DataFrame().style.to_html(**H),
    lambda m: m.DataFrame({"a": [1]}).style.pipe(lambda s, x: s.set_caption(x), "c").to_html(**H),
    lambda m: m.DataFrame({"a": [1]}).style.to_html(**H, doctype_html=True, bold_headers=True),
    lambda m: m.DataFrame({"a": [1.123456789]}).style.set_uuid("u")._repr_html_(),
    lambda m: mixed(m).style.to_latex(),
    lambda m: levels(m).style.to_latex(hrules=True, clines="all;data", caption="c", label="l"),
    lambda m: (
        mixed(m)
        .style.highlight_max(subset=["a"], props="font-weight: bold;")
        .to_latex(convert_css=True)
    ),
    lambda m: m.DataFrame({"a": [1]}).style.to_latex(environment="longtable", caption=("l", "s")),
    lambda m: m.DataFrame({"a": ["<b>&", "c_%$x$"]}).style.format(escape="latex-math").to_latex(),
    lambda m: levels(m).style.to_string(),
    lambda m: mixed(m).style.to_typst(),
]


@both
@pytest.mark.parametrize("build", TABLES)
def test_the_table_is_pandas_table(firepanda: ModuleType, build: Callable[[Any], Any]) -> None:
    """The same HTML, LaTeX, text or Typst as pandas writes."""
    import pandas as pd

    assert build(firepanda) == build(pd)


GRADIENTS: list[Callable[[Any], Any]] = [
    lambda m: mixed(m).style.background_gradient().to_html(**H),
    lambda m: mixed(m).style.text_gradient(cmap="viridis", axis=None).to_html(**H),
    lambda m: mixed(m).style.bar(cmap="viridis").to_html(**H),
    lambda m: small(m).style.background_gradient(gmap=[[1, 2], [3, 4]], axis=None).to_html(**H),
    lambda m: small(m).style.background_gradient(axis=1, low=0.2, high=0.3, vmin=0).to_html(**H),
]


@both
@pytest.mark.parametrize("build", GRADIENTS)
def test_a_gradient_is_pandas_gradient(firepanda: ModuleType, build: Callable[[Any], Any]) -> None:
    """Colors from a matplotlib map are the colors pandas picks."""
    pytest.importorskip("matplotlib")
    import pandas as pd

    assert build(firepanda) == build(pd)


def unplaced(text: str) -> str:
    """A message with the memory addresses of functions taken out."""
    return re.sub(r" at 0x[0-9a-f]+", "", text)


MISTAKES: list[Callable[[Any], Any]] = [
    lambda m: m.DataFrame({"a": [1]}).style.format(3),
    lambda m: m.DataFrame({"a": [1]}).style.map(lambda v: "red").to_html(),
    lambda m: m.DataFrame({"a": [1]}).style.apply(lambda c: "x").to_html(),
    lambda m: m.DataFrame({"a": [1]}).style.apply(lambda d: [[1]], axis=None).to_html(),
    lambda m: m.DataFrame({"a": [1]}).style.hide(axis=2),
    lambda m: m.DataFrame({"a": [1]}).style.set_caption(3),
    lambda m: m.DataFrame({"a": [1]}).style.bar(width=200),
    lambda m: m.DataFrame({"a": [1]}).style.bar(align="x").to_html(),
    lambda m: m.DataFrame({"a": [1]}).style.highlight_between(inclusive="x").to_html(),
    lambda m: m.DataFrame({"a": [1]}).style.to_latex(clines="x"),
    lambda m: m.DataFrame({"a": [1]}).style.relabel_index(["a", "b"]),
    lambda m: m.DataFrame({"a": ["x"]}).style.format(escape="x").to_html(),
    lambda m: small(m).style.highlight_between(left=[0, 3, 4]).to_html(),
    lambda m: m.DataFrame({"a": [1, 2]}, index=["x", "x"]).style.highlight_max().to_html(),
]


@both
@pytest.mark.parametrize("build", MISTAKES)
def test_a_mistake_is_pandas_mistake(firepanda: ModuleType, build: Callable[[Any], Any]) -> None:
    """The same class of error and the same message."""
    import pandas as pd

    with pytest.raises(Exception) as theirs:
        build(pd)
    with pytest.raises(Exception) as mine:
        build(firepanda)
    assert isinstance(mine.value, type(theirs.value))
    assert unplaced(str(mine.value)) == unplaced(str(theirs.value))


CSS = (
    "color: red; background-color: #ffff00; font-weight: bold; font-style: italic;"
    " border: 1px solid blue; text-align: center; vertical-align: top; number-format: 0.00;"
    " font-size: 14px; font-family: Arial, sans-serif; text-decoration: underline;"
    " white-space: normal"
)

WORKBOOKS: list[tuple[Callable[[Any], Any], dict[str, Any]]] = [
    (lambda m: mixed(m).style, {}),
    (lambda m: mixed(m).style.highlight_max(subset=["a", "b"]), {}),
    (lambda m: mixed(m).style.map(lambda v: CSS if isinstance(v, str) else "", subset=["s"]), {}),
    (
        lambda m: mixed(m).style.map(
            lambda v: "background: lightblue; border-bottom: 3px dashed #f00", subset=["b"]
        ),
        {},
    ),
    (
        lambda m: (
            levels(m)
            .style.map_index(lambda v: "font-weight: bold; color: green", level=0)
            .map_index(lambda v: "background-color: red", axis=1)
        ),
        {},
    ),
    (lambda m: levels(m).style.highlight_min(), {"merge_cells": False}),
    (
        lambda m: mixed(m).style.map(
            lambda v: "color: #f0a; border-top: thick double; font: bold 12pt serif"
        ),
        {"index": False},
    ),
]


def cells(path: Any) -> list[str]:
    """Each cell of a workbook with its value and the parts of its style a writer sets."""
    import openpyxl

    def side(edge: Any) -> Any:
        return getattr(edge, "style", None), getattr(getattr(edge, "color", None), "rgb", None)

    book = openpyxl.load_workbook(path)
    out = []
    for sheet in book.worksheets:
        out.append(str(sorted(str(span) for span in sheet.merged_cells.ranges)))
        for row in sheet.iter_rows():
            for cell in row:
                font, fill, edges, align = cell.font, cell.fill, cell.border, cell.alignment
                color = font.color.rgb if font.color else None
                out.append(
                    f"{cell.coordinate} {cell.value!r} {cell.number_format}"
                    f" {font.b} {font.i} {font.u} {font.sz} {font.name} {color}"
                    f" {fill.fill_type} {fill.fgColor.rgb if fill.fgColor else None}"
                    f" {[side(e) for e in (edges.top, edges.bottom, edges.left, edges.right)]}"
                    f" {align.horizontal} {align.vertical} {align.wrap_text}"
                )
    return out


@both
@pytest.mark.parametrize("engine", ["openpyxl", "xlsxwriter"])
@pytest.mark.parametrize("workbook", WORKBOOKS)
def test_a_styled_workbook_is_pandas_workbook(
    firepanda: ModuleType, tmp_path: Any, engine: str, workbook: Any
) -> None:
    """Every cell carries the value, font, fill, border and alignment pandas writes."""
    pytest.importorskip("openpyxl")
    pytest.importorskip(engine)
    import pandas as pd

    build, options = workbook
    build(pd).to_excel(tmp_path / "pd.xlsx", engine=engine, **options)
    build(firepanda).to_excel(tmp_path / "fp.xlsx", engine=engine, **options)
    assert cells(tmp_path / "fp.xlsx") == cells(tmp_path / "pd.xlsx")


@both
@pytest.mark.parametrize("workbook", WORKBOOKS)
def test_a_styled_spreadsheet_is_pandas_spreadsheet(
    firepanda: ModuleType, tmp_path: Any, workbook: Any
) -> None:
    """The OpenDocument styles and cells are the ones pandas writes."""
    pytest.importorskip("odf")
    import pandas as pd

    build, options = workbook
    written = []
    for module in (pd, firepanda):
        path = tmp_path / f"{module.__name__}.ods"
        build(module).to_excel(path, engine="odf", **options)
        with zipfile.ZipFile(path) as book:
            text = book.read("content.xml") + book.read("styles.xml")
        # odfpy declares a namespace for each of its modules loaded so far in the
        # process, which depends on what ran before, so the declarations are dropped.
        written.append(re.sub(rb' xmlns:\w+="[^"]*"', b"", text))
    assert written[0] == written[1]


def test_items_after_a_round_trip_are_the_labels(firepanda: ModuleType) -> None:
    """A column of a frame transposed twice hands back its labels, not their written text."""
    frame = firepanda.DataFrame({"a": [1, 5], "b": [2, 6]}).T.T
    assert list(frame["a"].items()) == [(0, 1), (1, 5)]


def test_no_jinja2_is_pandas_attribute_error(
    firepanda: ModuleType, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Without jinja2 the accessor is refused as pandas refuses it."""
    real = importlib.util.find_spec

    def hidden(name: str, *args: Any) -> Any:
        return None if name == "jinja2" else real(name, *args)

    monkeypatch.setattr(importlib.util, "find_spec", hidden)
    with pytest.raises(AttributeError, match=r"The '\.style' accessor requires jinja2"):
        _ = firepanda.DataFrame({"a": [1]}).style
